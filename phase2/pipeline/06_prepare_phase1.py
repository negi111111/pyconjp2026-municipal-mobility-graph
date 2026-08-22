"""Prepare the reproducible Phase 1 public-data graph tables.

This is the structured-data path for the current project.  It intentionally
does not read Wikipedia, call an LLM, or run GraphRAG.  All ranks and derived
facts are calculated locally from the pinned public datasets.
"""

from __future__ import annotations

import argparse
import io
import json
import math
import re
import sys
import zipfile
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import networkx as nx
import pandas as pd
import yaml
from _common import ENCODINGS, Ssdse
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
PROCESSED = ROOT / "data" / "processed"
CITY_CONFIG = ROOT / "configs" / "phase1_cities.yml"
FEATURE_CONFIG = ROOT / "configs" / "similarity_features.yml"

SSDSE_FILE = "SSDSE-A-2026.csv"
SSDSE_SOURCE_URL = "https://www.nstac.go.jp/files/SSDSE-A-2026.csv"
CENSUS_SOURCE_URL = "https://www.e-stat.go.jp/stat-search/database?statdisp_id=0003454524"
P05_SOURCE_URL = "https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-P05-v3_0.html"
C28_SOURCE_URL = "https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-C28-2021.html"
AIRPORT_PASSENGER_SOURCE_URL = "https://www.mlit.go.jp/koku/15_bf_000185.html"

CENSUS_FILES = {
    "32": ("table3_shimane.xlsx", "000032214257"),
    "33": ("table3_okayama.xlsx", "000032214260"),
    "34": ("table3_hiroshima.xlsx", "000032214263"),
    "35": ("table3_yamaguchi.xlsx", "000032214266"),
    "37": ("table3_kagawa.xlsx", "000032214272"),
    "38": ("table3_ehime.xlsx", "000032214275"),
}

AIRPORT_TYPE = {
    1: "拠点空港（会社管理）",
    2: "拠点空港（国管理）",
    3: "拠点空港（特定地方管理）",
    4: "地方管理空港",
    5: "その他の空港",
    6: "共用空港",
}


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _json_default(value: Any) -> Any:
    if value is pd.NA:
        return None
    if hasattr(value, "item"):
        return value.item()
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()
    if isinstance(value, set):
        return sorted(value)
    raise TypeError(f"cannot encode {type(value).__name__}")


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            value,
            ensure_ascii=False,
            indent=2,
            default=_json_default,
            allow_nan=False,
        ),
        encoding="utf-8",
    )


def load_city_master(path: Path = CITY_CONFIG) -> pd.DataFrame:
    config = yaml.safe_load(path.read_text(encoding="utf-8"))
    cities = pd.DataFrame(config["cities"])
    required = {"municipality_id", "municipality_code", "prefecture", "name"}
    missing = required.difference(cities.columns)
    if missing:
        raise ValueError(f"city master is missing columns: {sorted(missing)}")
    if len(cities) != 20:
        raise ValueError(f"Phase 1 requires exactly 20 cities; found {len(cities)}")
    for column in ("municipality_id", "municipality_code", "name"):
        if cities[column].duplicated().any():
            raise ValueError(f"city master contains duplicate {column}")
    if not cities["municipality_id"].str.fullmatch(r"R\d{5}").all():
        raise ValueError("municipality_id must use the SSDSE R##### form")
    if not cities["municipality_code"].astype(str).str.fullmatch(r"\d{5}").all():
        raise ValueError("municipality_code must be five digits")
    return cities.reset_index(drop=True)


def load_pinned_ssdse(path: Path) -> Ssdse:
    raw: pd.DataFrame | None = None
    encoding = "unknown"
    for candidate in ENCODINGS:
        try:
            raw = pd.read_csv(path, header=None, dtype=str, encoding=candidate)
            encoding = candidate
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        raise UnicodeError(f"could not decode pinned SSDSE file: {path}")
    if len(raw) < 4:
        raise ValueError(f"SSDSE file has no data rows: {path}")

    codes = raw.iloc[0].fillna("").astype(str).tolist()
    years = raw.iloc[1].fillna("").astype(str).tolist()
    labels = raw.iloc[2].fillna("").astype(str).tolist()
    data = raw.iloc[3:].reset_index(drop=True)
    data.columns = labels
    return Ssdse(path, codes, years, labels, data, encoding)


def _ssdse_column(year: str, label: str) -> str:
    return f"{year}_{label}" if year else label


def _numeric(series: pd.Series) -> pd.Series:
    cleaned = series.astype("string").str.strip().replace({"": pd.NA, "-": pd.NA})
    return pd.to_numeric(cleaned, errors="coerce")


def _ratio(numerator: pd.Series, denominator: pd.Series, scale: float = 1.0) -> pd.Series:
    numerator = pd.to_numeric(numerator, errors="coerce")
    denominator = pd.to_numeric(denominator, errors="coerce").where(lambda values: values > 0)
    return numerator / denominator * scale


def add_similarity_features(frame: pd.DataFrame) -> pd.DataFrame:
    result = frame.copy()
    population = result["2020_総人口"]
    children = result["2020_15歳未満人口"]
    households = result["2020_世帯数"]
    working_age = result["2020_15～64歳人口"]

    result["feature_population_density_per_km2"] = _ratio(population, result["2024_可住地面積"])
    result["feature_under15_share"] = _ratio(children, population)
    result["feature_working_age_share"] = _ratio(working_age, population)
    result["feature_over65_share"] = _ratio(result["2020_65歳以上人口"], population)
    result["feature_single_household_share"] = _ratio(result["2020_単独世帯数"], households)
    result["feature_household_size"] = _ratio(population, households)

    result["feature_hospitals_per_10k"] = _ratio(result["2023_一般病院数"], population, 10_000)
    result["feature_clinics_per_10k"] = _ratio(result["2023_一般診療所数"], population, 10_000)
    result["feature_dental_clinics_per_10k"] = _ratio(
        result["2023_歯科診療所数"], population, 10_000
    )
    result["feature_doctors_per_10k"] = _ratio(result["2022_医師数"], population, 10_000)
    result["feature_dentists_per_10k"] = _ratio(result["2022_歯科医師数"], population, 10_000)
    result["feature_pharmacists_per_10k"] = _ratio(result["2022_薬剤師数"], population, 10_000)

    result["feature_elementary_students_per_teacher"] = _ratio(
        result["2024_小学校児童数"], result["2024_小学校教員数"]
    )
    result["feature_junior_students_per_teacher"] = _ratio(
        result["2024_中学校生徒数"], result["2024_中学校教員数"]
    )
    result["feature_schools_per_10k_children"] = _ratio(
        result["2024_小学校数"] + result["2024_中学校数"], children, 10_000
    )
    result["feature_high_school_students_per_1k"] = _ratio(
        result["2024_高等学校生徒数"], population, 1_000
    )
    result["feature_libraries_per_10k"] = _ratio(result["2021_図書館数"], population, 10_000)

    employed = result["2020_就業者数"]
    unemployed = result["2020_完全失業者数"]
    result["feature_employment_rate"] = _ratio(employed, working_age)
    result["feature_unemployment_rate"] = _ratio(unemployed, employed + unemployed)
    result["feature_private_establishments_per_1k"] = _ratio(
        result["2021_事業所数（民営）"], population, 1_000
    )
    result["feature_private_workers_per_1k"] = _ratio(
        result["2021_従業者数（民営）"], population, 1_000
    )
    result["feature_tertiary_sector_share"] = _ratio(result["2020_第3次産業就業者数"], employed)
    result["feature_retail_stores_per_1k"] = _ratio(result["2021_小売店数"], population, 1_000)

    result["feature_elderly_single_household_share"] = _ratio(
        result["2020_65歳以上世帯員の単独世帯数"], households
    )
    result["feature_childcare_facilities_per_1k_children"] = _ratio(
        result["2023_保育所等数（基本票）"], children, 1_000
    )
    result["feature_welfare_establishments_per_1k"] = _ratio(
        result["2021_事業所数（民営）（医療、福祉）"], population, 1_000
    )
    result["feature_municipal_welfare_spending_per_capita"] = _ratio(
        result["2022_民生費（市町村財政）"], population
    )
    return result


def load_city_halls(cities: pd.DataFrame, archive: Path) -> pd.DataFrame:
    by_name: dict[str, list[dict[str, Any]]] = defaultdict(list)
    with zipfile.ZipFile(archive) as outer:
        for prefix in sorted(cities["municipality_code"].str[:2].unique()):
            inner_name = f"P05-22_{prefix}_GML.zip"
            with zipfile.ZipFile(io.BytesIO(outer.read(inner_name))) as inner:
                payload = json.loads(inner.read(f"P05-22_{prefix}.geojson"))
            for feature in payload["features"]:
                properties = feature["properties"]
                longitude, latitude = feature["geometry"]["coordinates"]
                by_name[str(properties["P05_003"])].append(
                    {
                        "facility_admin_code": str(properties["P05_001"]),
                        "facility_class_code": str(properties["P05_002"]),
                        "representative_point_name": str(properties["P05_003"]),
                        "representative_point_address": str(properties.get("P05_004") or ""),
                        "latitude": float(latitude),
                        "longitude": float(longitude),
                    }
                )

    records: list[dict[str, Any]] = []
    for city in cities.to_dict(orient="records"):
        expected_name = f"{city['name']}役所"
        matches = [
            row for row in by_name.get(expected_name, []) if row["facility_class_code"] == "1"
        ]
        if len(matches) != 1:
            raise ValueError(
                f"P05 main-office match for {city['prefecture']}{city['name']} "
                f"is not unique: {len(matches)}"
            )
        records.append(
            {
                "municipality_id": city["municipality_id"],
                "municipality_code": city["municipality_code"],
                **matches[0],
                "coordinate_role": "municipal_main_office",
                "coordinate_reference_year": 2022,
                "coordinate_source_url": P05_SOURCE_URL,
            }
        )
    return pd.DataFrame(records)


def prepare_municipalities(
    cities: pd.DataFrame,
    raw_dir: Path,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, dict[str, dict[str, str]]]:
    ssdse = load_pinned_ssdse(raw_dir / SSDSE_FILE)
    metric_columns: list[str] = []
    catalog_rows: list[dict[str, Any]] = []
    rename: dict[str, str] = {}
    for code, year, label in zip(ssdse.codes[3:], ssdse.years[3:], ssdse.labels[3:], strict=True):
        output = _ssdse_column(year, label)
        if output in metric_columns:
            raise ValueError(f"duplicate SSDSE output column: {output}")
        metric_columns.append(output)
        rename[label] = output
        catalog_rows.append(
            {
                "indicator_code": code,
                "reference_year": year,
                "label": label,
                "column": output,
                "source_url": SSDSE_SOURCE_URL,
            }
        )

    data = ssdse.data.copy()
    canonical: dict[str, dict[str, str]] = {}
    for row in data[["地域コード", "都道府県", "市区町村"]].to_dict(orient="records"):
        municipality_id = str(row["地域コード"])
        if not re.fullmatch(r"R\d{5}", municipality_id):
            continue
        canonical[municipality_id[1:]] = {
            "municipality_id": municipality_id,
            "name": str(row["市区町村"]),
            "prefecture": str(row["都道府県"]),
        }

    data = data.set_index("地域コード", drop=False)
    records: list[dict[str, Any]] = []
    long_rows: list[dict[str, Any]] = []
    for city in cities.to_dict(orient="records"):
        municipality_id = city["municipality_id"]
        if municipality_id not in data.index:
            raise ValueError(f"SSDSE does not contain {municipality_id}")
        source = data.loc[municipality_id]
        if source["都道府県"] != city["prefecture"] or source["市区町村"] != city["name"]:
            raise ValueError(f"SSDSE label mismatch for {municipality_id}")
        record: dict[str, Any] = dict(city)
        for source_label, output in rename.items():
            value = pd.to_numeric(
                pd.Series([source[source_label]]).replace({"-": pd.NA}), errors="coerce"
            ).iloc[0]
            record[output] = value
        records.append(record)
        for item in catalog_rows:
            long_rows.append(
                {
                    "municipality_id": municipality_id,
                    "municipality_code": city["municipality_code"],
                    "prefecture": city["prefecture"],
                    "name": city["name"],
                    "indicator_code": item["indicator_code"],
                    "reference_year": item["reference_year"],
                    "label": item["label"],
                    "value": record[item["column"]],
                    "source_url": SSDSE_SOURCE_URL,
                }
            )

    municipalities = pd.DataFrame(records)
    city_halls = load_city_halls(cities, raw_dir / "geography" / "P05-22_GML.zip")
    municipalities = municipalities.merge(
        city_halls.drop(columns=["municipality_code"]),
        on="municipality_id",
        how="left",
        validate="one_to_one",
    )
    municipalities = add_similarity_features(municipalities)
    core = [
        "municipality_id",
        "municipality_code",
        "prefecture",
        "name",
        "latitude",
        "longitude",
        "representative_point_name",
        "representative_point_address",
        "coordinate_role",
        "coordinate_reference_year",
        "coordinate_source_url",
        "facility_admin_code",
        "facility_class_code",
    ]
    municipalities = municipalities[core + [c for c in municipalities if c not in core]]
    return (
        municipalities,
        pd.DataFrame(long_rows),
        pd.DataFrame(catalog_rows),
        canonical,
    )


def _area_code(value: Any) -> str | None:
    match = re.match(r"^(\d{5})_", str(value or ""))
    return match.group(1) if match else None


def _census_number(value: Any, symbols: Counter[str]) -> int | None:
    if value is None or str(value).strip() == "":
        symbols["blank"] += 1
        return None
    if isinstance(value, bool):
        raise ValueError(f"unexpected boolean Census value: {value}")
    if isinstance(value, (int, float)):
        if isinstance(value, float) and math.isnan(value):
            symbols["nan"] += 1
            return None
        return int(value)
    text = str(value).strip()
    if text in {"-", "—", "–"}:
        symbols["dash_as_zero"] += 1
        return 0
    symbols[f"unparsed:{text}"] += 1
    return None


def prepare_flows(
    cities: pd.DataFrame,
    canonical: dict[str, dict[str, str]],
    raw_dir: Path,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame, dict[str, Any]]:
    seed_codes = set(cities["municipality_code"].astype(str))
    seed_id = dict(zip(cities["municipality_code"], cities["municipality_id"], strict=True))
    source_prefixes = set(CENSUS_FILES)
    symbols: Counter[str] = Counter()
    induced: list[dict[str, Any]] = []
    diagnostics: list[dict[str, Any]] = []
    discrepancy_examples: list[dict[str, Any]] = []
    scanned_rows = 0

    for _prefix, (filename, stat_inf_id) in CENSUS_FILES.items():
        path = raw_dir / "census_od" / filename
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        for row in sheet.iter_rows(min_row=11, max_col=14, values_only=True):
            scanned_rows += 1
            if row[0] != "0_総数":
                continue
            source_code = _area_code(row[3])
            target_code = _area_code(row[8])
            if source_code not in canonical or target_code not in canonical:
                continue
            if source_code == target_code:
                continue

            source_prefecture = canonical[source_code]["prefecture"]
            target_prefecture = canonical[target_code]["prefecture"]
            source_is_seed = source_code in seed_codes
            target_is_seed = target_code in seed_codes
            target_in_source_region = target_code[:2] in source_prefixes
            if not (
                (source_is_seed and target_is_seed)
                or ((source_is_seed or target_is_seed) and target_in_source_region)
            ):
                continue

            raw_total = _census_number(row[9], symbols)
            commuters = _census_number(row[11], symbols)
            students_15plus = _census_number(row[12], symbols)
            students = _census_number(row[13], symbols)
            if commuters is None or students is None or raw_total is None:
                continue
            computed_total = commuters + students
            if computed_total != raw_total and len(discrepancy_examples) < 20:
                discrepancy_examples.append(
                    {
                        "source": source_code,
                        "target": target_code,
                        "raw_total": raw_total,
                        "computed_total": computed_total,
                    }
                )
            if computed_total <= 0:
                continue

            fact = {
                "source_code": source_code,
                "target_code": target_code,
                "source_name": canonical[source_code]["name"],
                "target_name": canonical[target_code]["name"],
                "source_prefecture": source_prefecture,
                "target_prefecture": target_prefecture,
                "commuters": commuters,
                "students_15plus": students_15plus,
                "students": students,
                "total": computed_total,
                "raw_total": raw_total,
                "year": 2020,
                "cross_prefecture": source_prefecture != target_prefecture,
                "source_stat_inf_id": stat_inf_id,
                "source_url": CENSUS_SOURCE_URL,
            }
            if source_is_seed and target_is_seed:
                induced.append(
                    {
                        "source": seed_id[source_code],
                        "target": seed_id[target_code],
                        **fact,
                    }
                )
            if source_is_seed != target_is_seed:
                diagnostics.append(fact)
        workbook.close()
        print(f"[OD] {filename}: parsed")

    flows = pd.DataFrame(induced)
    if flows.empty:
        raise ValueError("no induced Phase 1 Census OD rows were found")
    flows = flows.sort_values(["source", "target"], kind="stable").reset_index(drop=True)
    duplicate_count = int(flows.duplicated(["source", "target"]).sum())
    commute = flows.loc[flows["commuters"] > 0].copy()
    study = flows.loc[flows["students"] > 0].copy()

    diagnostics_frame = pd.DataFrame(diagnostics)
    review_rows: list[dict[str, Any]] = []
    if not diagnostics_frame.empty:
        pair_totals: dict[tuple[str, str], dict[str, Any]] = {}
        for row in diagnostics_frame.to_dict(orient="records"):
            if row["source_code"] in seed_codes:
                seed, candidate = row["source_code"], row["target_code"]
                direction = "seed_to_candidate"
            else:
                seed, candidate = row["target_code"], row["source_code"]
                direction = "candidate_to_seed"
            entry = pair_totals.setdefault(
                (candidate, seed),
                {"seed_to_candidate": 0, "candidate_to_seed": 0},
            )
            entry[direction] += int(row["total"])

        candidate_totals: dict[str, dict[str, Any]] = {}
        for (candidate, seed), values in pair_totals.items():
            total = values["seed_to_candidate"] + values["candidate_to_seed"]
            entry = candidate_totals.setdefault(
                candidate,
                {
                    "candidate_code": candidate,
                    "candidate_name": canonical[candidate]["name"],
                    "candidate_prefecture": canonical[candidate]["prefecture"],
                    "bidirectional_flow_with_seed": 0,
                    "strongest_seed_code": seed,
                    "strongest_seed_name": canonical[seed]["name"],
                    "strongest_seed_bidirectional_flow": -1,
                },
            )
            entry["bidirectional_flow_with_seed"] += total
            if total > entry["strongest_seed_bidirectional_flow"]:
                entry["strongest_seed_code"] = seed
                entry["strongest_seed_name"] = canonical[seed]["name"]
                entry["strongest_seed_bidirectional_flow"] = total
        review_rows = sorted(
            candidate_totals.values(),
            key=lambda row: (-row["bidirectional_flow_with_seed"], row["candidate_code"]),
        )
        for rank, row in enumerate(review_rows, start=1):
            row["rank"] = rank
            row["decision"] = "review_only_seed_unchanged"

    diagnostics_report = {
        "scanned_rows": scanned_rows,
        "induced_positive_flow_rows": len(flows),
        "commute_rows": len(commute),
        "study_rows": len(study),
        "cross_prefecture_rows": int(flows["cross_prefecture"].sum()),
        "duplicate_source_target_rows": duplicate_count,
        "total_reconciliation_failures": len(discrepancy_examples),
        "total_reconciliation_examples": discrepancy_examples,
        "source_symbol_handling": dict(sorted(symbols.items())),
        "self_loops_excluded": True,
        "self_loop_reason": (
            "Census footnote assigns unknown/foreign destinations to residence; "
            "self-loops are not reliable intercity movements."
        ),
        "seed_membership_review_candidate_count": len(review_rows),
    }
    return flows, commute, study, pd.DataFrame(review_rows), diagnostics_report


def _airport_name_key(name: str) -> str:
    compact = re.sub(r"[\s\u3000]+", "", str(name))
    if "奈多地区" in compact:
        return "福岡奈多"
    compact = re.sub(r"[（(].*$", "", compact)
    compact = re.sub(r"(空港|飛行場)$", "", compact)
    return {"新石垣": "石垣"}.get(compact, compact)


def _airport_scale(passengers: int | None) -> tuple[float | None, str]:
    if passengers is None:
        return None, "unknown"
    if passengers >= 10_000_000:
        return 4.0, "major_10m_plus"
    if passengers >= 3_000_000:
        return 3.0, "large_3m_plus"
    if passengers >= 1_000_000:
        return 2.0, "regional_1m_plus"
    return 1.0, "local_under_1m"


def load_airport_passengers(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook["空港管理状況調書"]
    rows = list(sheet.iter_rows(values_only=True))
    result: dict[str, dict[str, Any]] = {}
    for index, row in enumerate(rows):
        if len(row) <= 3 or row[2] != "空港名 ：" or index + 18 >= len(rows):
            continue
        annual = rows[index + 18]
        if annual[2] != "暦年 計":
            continue
        key = _airport_name_key(str(row[3]))
        if key.endswith("計") or key.startswith("全空港計"):
            continue
        values = {
            "annual_passengers_year": 2025,
            "international_passengers": int(annual[9] or 0),
            "domestic_passengers": int(annual[12] or 0),
            "annual_passengers": int(annual[13] or 0),
        }
        if key in result and result[key] != values:
            # The workbook also contains same-named heliports.  Keep the first
            # passenger-airport block; C28 never joins to a heliport here.
            continue
        result[key] = values
    workbook.close()
    return result


def prepare_airports(raw_dir: Path) -> tuple[pd.DataFrame, dict[str, Any]]:
    passengers = load_airport_passengers(raw_dir / "airports" / "airport_management_2025.xlsx")
    archive = raw_dir / "airports" / "C28-21_GML.zip"
    with zipfile.ZipFile(archive) as source:
        airport_data = json.loads(source.read("UTF-8/C28-21_Airport.geojson"))
        point_data = json.loads(source.read("UTF-8/C28-21_AirportReferencePoint.geojson"))
    points = {
        feature["properties"]["C28_000"]: feature["geometry"]["coordinates"]
        for feature in point_data["features"]
    }

    by_reference: dict[str, dict[str, Any]] = {}
    for feature in airport_data["features"]:
        properties = feature["properties"]
        regular_code = int(properties["C28_011"])
        if properties["C28_004"] != "供用中" or regular_code not in {1, 2}:
            continue
        reference = str(properties["C28_101"]).lstrip("#")
        if reference not in points:
            raise ValueError(f"C28 airport reference point is missing: {reference}")
        longitude, latitude = points[reference]
        name = str(properties["C28_005"])
        passenger = passengers.get(_airport_name_key(name))
        scale, scale_label = _airport_scale(
            passenger["annual_passengers"] if passenger is not None else None
        )
        candidate = {
            "airport_id": f"AP_{reference}",
            "name": name,
            "latitude": float(latitude),
            "longitude": float(longitude),
            "coordinate_role": "official_airport_reference_point",
            "airport_type_code": int(properties["C28_003"]),
            "airport_type": AIRPORT_TYPE[int(properties["C28_003"])],
            "scheduled_service": True,
            "scheduled_service_type": "jet" if regular_code == 1 else "propeller",
            "scheduled_service_code": regular_code,
            "annual_passengers": (
                passenger["annual_passengers"] if passenger is not None else pd.NA
            ),
            "domestic_passengers": (
                passenger["domestic_passengers"] if passenger is not None else pd.NA
            ),
            "international_passengers": (
                passenger["international_passengers"] if passenger is not None else pd.NA
            ),
            "annual_passengers_year": 2025 if passenger is not None else pd.NA,
            "airport_scale": scale,
            "airport_scale_label": scale_label,
            "source_url": C28_SOURCE_URL,
            "passenger_source_url": AIRPORT_PASSENGER_SOURCE_URL,
            "c28_reference_id": reference,
        }
        previous = by_reference.get(reference)
        if previous is None or candidate["name"] < previous["name"]:
            by_reference[reference] = candidate

    airports = pd.DataFrame(by_reference.values()).sort_values("airport_id").reset_index(drop=True)
    diagnostics = {
        "operating_scheduled_airports": len(airports),
        "passenger_match_count": int(airports["annual_passengers"].notna().sum()),
        "passenger_match_rate": float(airports["annual_passengers"].notna().mean()),
        "unmatched_airports": airports.loc[airports["annual_passengers"].isna(), "name"].tolist(),
        "regular_service_code": {
            "1": "jet scheduled service",
            "2": "propeller scheduled service",
        },
        "passenger_total_definition": (
            "MLIT calendar-year total; international subtotal includes transit passengers"
        ),
    }
    return airports, diagnostics


def haversine_km(
    latitude_a: float,
    longitude_a: float,
    latitude_b: float,
    longitude_b: float,
) -> float:
    radius_km = 6_371.0088
    lat_a, lon_a, lat_b, lon_b = map(
        math.radians, (latitude_a, longitude_a, latitude_b, longitude_b)
    )
    delta_lat = lat_b - lat_a
    delta_lon = lon_b - lon_a
    value = (
        math.sin(delta_lat / 2) ** 2
        + math.cos(lat_a) * math.cos(lat_b) * math.sin(delta_lon / 2) ** 2
    )
    return 2 * radius_km * math.asin(math.sqrt(value))


def prepare_airport_access(
    municipalities: pd.DataFrame,
    airports: pd.DataFrame,
) -> pd.DataFrame:
    records: list[dict[str, Any]] = []
    airport_rows = airports.to_dict(orient="records")
    for city in municipalities.to_dict(orient="records"):
        ranked: list[tuple[float, dict[str, Any]]] = []
        for airport in airport_rows:
            distance = haversine_km(
                float(city["latitude"]),
                float(city["longitude"]),
                float(airport["latitude"]),
                float(airport["longitude"]),
            )
            ranked.append((distance, airport))
        distance, airport = min(ranked, key=lambda item: (item[0], item[1]["airport_id"]))
        records.append(
            {
                "source": city["municipality_id"],
                "target": airport["airport_id"],
                "source_name": city["name"],
                "target_name": airport["name"],
                "distance_km": distance,
                "distance_method": "Haversine_WGS84_mean_earth_radius",
                "rank": 1,
                "airport_scale": airport["airport_scale"],
                "airport_scale_label": airport["airport_scale_label"],
                "annual_passengers": airport["annual_passengers"],
                "annual_passengers_year": airport["annual_passengers_year"],
                "municipality_coordinate_role": city["coordinate_role"],
                "airport_coordinate_role": airport["coordinate_role"],
            }
        )
    return pd.DataFrame(records).sort_values("source").reset_index(drop=True)


def load_feature_specs(path: Path = FEATURE_CONFIG) -> tuple[list[Any], str]:
    from app.analysis import FeatureSpec

    config = yaml.safe_load(path.read_text(encoding="utf-8"))
    allowed = {"name", "category", "per_capita", "denominator", "per_capita_scale", "log1p"}
    specs = [
        FeatureSpec(**{key: value for key, value in item.items() if key in allowed})
        for item in config["features"]
    ]
    return specs, str(config["population_column"])


def build_quality_report(
    municipalities: pd.DataFrame,
    ssdse_long: pd.DataFrame,
    flows: pd.DataFrame,
    airport_access: pd.DataFrame,
    airports: pd.DataFrame,
    graph: nx.MultiDiGraph,
    flow_diagnostics: dict[str, Any],
    airport_diagnostics: dict[str, Any],
) -> dict[str, Any]:
    from app.graph import graph_facts

    checks: list[dict[str, Any]] = []

    def check(name: str, passed: bool, observed: Any, expected: Any) -> None:
        checks.append(
            {
                "name": name,
                "status": "pass" if passed else "fail",
                "severity": "critical",
                "observed": observed,
                "expected": expected,
            }
        )

    check("municipality_count", len(municipalities) == 20, len(municipalities), 20)
    check(
        "municipality_id_unique",
        municipalities["municipality_id"].is_unique,
        municipalities["municipality_id"].nunique(),
        20,
    )
    coordinate_complete = municipalities[["latitude", "longitude"]].notna().all().all()
    check("municipality_coordinates_complete", bool(coordinate_complete), coordinate_complete, True)
    check(
        "main_office_rule",
        bool((municipalities["facility_class_code"] == "1").all()),
        municipalities["facility_class_code"].value_counts().to_dict(),
        {"1": 20},
    )
    metric_counts = ssdse_long.groupby("municipality_id")["indicator_code"].nunique()
    check(
        "ssdse_metric_coverage",
        bool((metric_counts == 125).all()),
        {"min": int(metric_counts.min()), "max": int(metric_counts.max())},
        {"min": 125, "max": 125},
    )
    check(
        "flow_endpoint_integrity",
        set(flows["source"])
        .union(flows["target"])
        .issubset(set(municipalities["municipality_id"])),
        len(set(flows["source"]).union(flows["target"])),
        "all endpoints in 20-city master",
    )
    check(
        "flow_source_target_unique",
        flow_diagnostics["duplicate_source_target_rows"] == 0,
        flow_diagnostics["duplicate_source_target_rows"],
        0,
    )
    check(
        "flow_total_reconciliation",
        flow_diagnostics["total_reconciliation_failures"] == 0,
        flow_diagnostics["total_reconciliation_failures"],
        0,
    )
    check("flow_no_self_loops", bool((flows["source"] != flows["target"]).all()), True, True)
    check(
        "cross_prefecture_flows_retained",
        bool(flows["cross_prefecture"].any()),
        int(flows["cross_prefecture"].sum()),
        "> 0",
    )
    check(
        "airport_passenger_join",
        airport_diagnostics["passenger_match_rate"] >= 0.9,
        airport_diagnostics["passenger_match_rate"],
        ">= 0.9",
    )
    check(
        "nearest_airport_coverage",
        len(airport_access) == 20 and airport_access["source"].is_unique,
        len(airport_access),
        20,
    )
    check(
        "nearest_airport_rank",
        bool((airport_access["rank"] == 1).all()),
        airport_access["rank"].value_counts().to_dict(),
        {1: 20},
    )
    facts = graph_facts(graph)
    check(
        "graph_node_layers",
        facts["node_counts"].get("Municipality") == 20
        and facts["node_counts"].get("Prefecture") == 6
        and facts["node_counts"].get("Airport") == len(airports),
        facts["node_counts"],
        {"Municipality": 20, "Prefecture": 6, "Airport": len(airports)},
    )
    check(
        "administrative_edges_not_flow_edges",
        facts["relation_counts"].get("IS_IN") == 20
        and facts["relation_counts"].get("ACCESS_TO") == 20,
        facts["relation_counts"],
        "20 IS_IN and 20 ACCESS_TO; flow relations separate",
    )
    failures = [item for item in checks if item["status"] == "fail"]
    return {
        "generated_at": utc_now(),
        "overall_status": "pass" if not failures else "fail",
        "dataset_grain": {
            "municipalities": "one row per fixed Phase 1 municipality",
            "flow_edges": "one directed source-target municipality pair for Census 2020",
            "airport_access": "one nearest scheduled-service airport per municipality",
        },
        "checks": checks,
        "failed_check_count": len(failures),
        "flow_diagnostics": flow_diagnostics,
        "airport_diagnostics": airport_diagnostics,
        "graph_facts": facts,
        "limitations": [
            "The 20-city flow graph is an induced subgraph and omits flows to outside cities.",
            "Census self-loops are excluded because unknown/foreign destinations "
            "are allocated to residence.",
            "Airport distance is straight-line Haversine distance, not road or rail travel time.",
            "C28 airport coordinates are official reference points, not terminal entrances.",
        ],
    }


def core_analysis_snapshot(
    municipalities: pd.DataFrame,
    graph: nx.MultiDiGraph,
) -> dict[str, Any]:
    from app.analysis import (
        city_similarity,
        compute_flow_centrality,
        detect_flow_communities,
        find_similar_cities_with_better_airport,
        rank_airport_access,
        strongest_bidirectional_flow,
    )

    specs, population_column = load_feature_specs()
    medical_specs = [spec for spec in specs if spec.category == "medical"]
    return {
        "generated_at": utc_now(),
        "strongest_flow_from_hiroshima": strongest_bidirectional_flow(graph, "R34100"),
        "flow_hubs": compute_flow_centrality(graph),
        "flow_communities_total": detect_flow_communities(graph, normalization="total"),
        "flow_communities_share": detect_flow_communities(graph, normalization="flow_share"),
        "overall_similarity_to_hiroshima": city_similarity(
            municipalities,
            "R34100",
            feature_specs=specs,
            population_column=population_column,
            top_n=10,
        ),
        "medical_similarity_to_hiroshima": city_similarity(
            municipalities,
            "R34100",
            feature_specs=medical_specs,
            population_column=population_column,
            top_n=10,
        ),
        "airport_access": rank_airport_access(graph, max_access_rank=1),
        "similar_with_better_airport": find_similar_cities_with_better_airport(
            municipalities,
            graph,
            "R34100",
            feature_specs=specs,
            population_column=population_column,
            max_access_rank=1,
            require_nearer=True,
            top_n=10,
        ),
    }


def run(raw_dir: Path = RAW, output_dir: Path = PROCESSED) -> dict[str, Any]:
    from app.graph import build_mobility_graph

    cities = load_city_master()
    municipalities, ssdse_long, metric_catalog, canonical = prepare_municipalities(cities, raw_dir)
    flows, commute, study, seed_review, flow_diagnostics = prepare_flows(cities, canonical, raw_dir)
    airports, airport_diagnostics = prepare_airports(raw_dir)
    airport_access = prepare_airport_access(municipalities, airports)
    graph = build_mobility_graph(
        municipalities,
        flow_edges=flows,
        commute_edges=commute,
        study_edges=study,
        airports=airports,
        airport_access=airport_access,
    )
    quality = build_quality_report(
        municipalities,
        ssdse_long,
        flows,
        airport_access,
        airports,
        graph,
        flow_diagnostics,
        airport_diagnostics,
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    write_json(output_dir / "quality_report.json", quality)
    if quality["overall_status"] != "pass":
        failed = [item["name"] for item in quality["checks"] if item["status"] == "fail"]
        raise ValueError(f"Phase 1 data quality checks failed: {', '.join(failed)}")

    tables = {
        "municipalities.parquet": municipalities,
        "ssdse_metrics.parquet": ssdse_long,
        "commute_edges.parquet": commute,
        "study_edges.parquet": study,
        "flow_edges.parquet": flows,
        "airports.parquet": airports,
        "airport_access.parquet": airport_access,
    }
    for filename, frame in tables.items():
        frame.to_parquet(output_dir / filename, index=False)
        print(f"[write] {output_dir / filename} ({len(frame)} rows)")
    metric_catalog.to_csv(output_dir / "ssdse_metric_catalog.csv", index=False)
    seed_review.to_csv(output_dir / "seed_membership_review.csv", index=False)
    municipalities.to_csv(output_dir / "municipalities.csv", index=False)
    flows.to_csv(output_dir / "flow_edges.csv", index=False)

    write_json(output_dir / "knowledge_graph.json", nx.node_link_data(graph))
    analyses = core_analysis_snapshot(municipalities, graph)
    write_json(output_dir / "phase1_analysis.json", analyses)
    report = {
        "generated_at": utc_now(),
        "pipeline": "pipeline/06_prepare_phase1.py",
        "phase": 1,
        "seed_city_count": 20,
        "outputs": {
            filename: {"rows": len(frame), "columns": len(frame.columns)}
            for filename, frame in tables.items()
        },
        "quality_report": "data/processed/quality_report.json",
        "analysis_snapshot": "data/processed/phase1_analysis.json",
        "wikipedia_used": False,
        "graphrag_used": False,
        "llm_generated_facts": False,
    }
    write_json(output_dir / "processing_report.json", report)
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=RAW)
    parser.add_argument("--output-dir", type=Path, default=PROCESSED)
    args = parser.parse_args()
    try:
        report = run(args.raw_dir, args.output_dir)
    except Exception as exc:
        print(f"[error] {exc}", file=sys.stderr)
        return 1
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
