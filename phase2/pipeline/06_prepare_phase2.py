"""Prepare the national analysis universe and the approved Phase 2 search set."""

from __future__ import annotations

import argparse
import importlib.util
import io
import json
import math
import re
import sys
import zipfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import networkx as nx
import pandas as pd
import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
PROCESSED = ROOT / "data" / "processed"
CITY_CONFIG = ROOT / "configs" / "phase2_cities.yml"
FEATURE_CONFIG = ROOT / "configs" / "similarity_features.yml"
ROLE_CONFIG = ROOT / "configs" / "phase2_city_roles.yml"
OD_DIR = RAW / "census_od_national"
P05_ARCHIVE = RAW / "geography" / "P05-22_GML.zip"
C28_ARCHIVE = RAW / "airports" / "C28-21_GML.zip"
SSDSE_FILE = RAW / "SSDSE-A-2026.csv"

sys.path.insert(0, str(ROOT / "pipeline"))
sys.path.insert(0, str(ROOT / "scripts"))

from _common import ENCODINGS  # noqa: E402
from download_national_selection_od import PREFECTURE_FILES  # noqa: E402


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _json_default(value: Any) -> Any:
    if value is pd.NA:
        return None
    if hasattr(value, "item"):
        return value.item()
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.isoformat()
    raise TypeError(f"cannot encode {type(value).__name__}")


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(
            value,
            ensure_ascii=False,
            indent=2,
            default=_json_default,
            allow_nan=False,
        )
        + "\n",
        encoding="utf-8",
    )


def _load_phase1_module():
    path = ROOT / "pipeline" / "06_prepare_phase1.py"
    spec = importlib.util.spec_from_file_location("phase1_prepare", path)
    if spec is None or spec.loader is None:
        raise ImportError(f"cannot load Phase 1 preparation module: {path}")
    module = importlib.util.module_from_spec(spec)
    sys.modules["phase1_prepare"] = module
    spec.loader.exec_module(module)
    return module


def _numeric(value: Any) -> int | None:
    if value is None or str(value).strip() in {"", "-", "—", "–"}:
        return 0
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _area_code(value: Any) -> str | None:
    match = re.match(r"^(\d{5})_", str(value or "").strip())
    return match.group(1) if match else None


def _normalize_tokyo_code(code: str | None) -> str | None:
    if code and code.startswith("131"):
        return "13000"
    return code


def _haversine_km(
    latitude_a: float,
    longitude_a: float,
    latitude_b: float,
    longitude_b: float,
) -> float:
    radius_km = 6_371.0088
    lat_a, lon_a, lat_b, lon_b = map(
        math.radians, (latitude_a, longitude_a, latitude_b, longitude_b)
    )
    delta_lat = lat_b - lat_a
    delta_lon = lon_b - lon_a
    value = (
        math.sin(delta_lat / 2) ** 2
        + math.cos(lat_a) * math.cos(lat_b) * math.sin(delta_lon / 2) ** 2
    )
    return 2 * radius_km * math.asin(math.sqrt(value))


def load_city_config() -> tuple[pd.DataFrame, dict[str, Any]]:
    config = yaml.safe_load(CITY_CONFIG.read_text(encoding="utf-8"))
    cities = pd.DataFrame(config["cities"])
    if len(cities) != 100:
        raise ValueError(f"Phase 2 requires exactly 100 search cities; found {len(cities)}")
    if cities["municipality_code"].astype(str).duplicated().any():
        raise ValueError("Phase 2 search city codes are not unique")
    return cities, config


def load_ssdse() -> tuple[pd.DataFrame, dict[str, dict[str, Any]], list[str]]:
    raw: pd.DataFrame | None = None
    for encoding in ENCODINGS:
        try:
            raw = pd.read_csv(SSDSE_FILE, header=None, dtype=str, encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        raise UnicodeError(f"cannot decode {SSDSE_FILE}")
    years = raw.iloc[1].fillna("").astype(str).tolist()
    labels = raw.iloc[2].fillna("").astype(str).tolist()
    columns = [
        label if index < 3 else (f"{year}_{label}" if year and label else label)
        for index, (year, label) in enumerate(zip(years, labels, strict=True))
    ]
    data = raw.iloc[3:].reset_index(drop=True)
    data.columns = columns
    data["地域コード"] = data["地域コード"].astype(str).str.strip()
    data["都道府県"] = data["都道府県"].astype(str).str.strip()
    data["市区町村"] = data["市区町村"].astype(str).str.strip()
    valid = data["地域コード"].str.fullmatch(r"R\d{5}")
    data = data.loc[valid].copy()
    metadata = ["地域コード", "都道府県", "市区町村"]
    metrics = [column for column in data.columns if column not in metadata]
    canonical = {
        str(row["地域コード"])[1:]: {
            "municipality_id": str(row["地域コード"]),
            "municipality_code": str(row["地域コード"])[1:],
            "prefecture": str(row["都道府県"]),
            "name": str(row["市区町村"]),
        }
        for row in data[metadata].to_dict(orient="records")
    }
    return data, canonical, metrics


def aggregate_tokyo_wards(
    data: pd.DataFrame,
    canonical: dict[str, dict[str, Any]],
    metrics: list[str],
) -> dict[str, Any]:
    ward_mask = (
        (data["都道府県"] == "東京都")
        & data["地域コード"].str[1:].str.startswith("131")
        & data["市区町村"].str.endswith("区")
    )
    wards = data.loc[ward_mask]
    if len(wards) != 23:
        raise ValueError(f"Tokyo ward aggregation expected 23 wards; found {len(wards)}")
    record: dict[str, Any] = {
        "地域コード": "R13000",
        "都道府県": "東京都",
        "市区町村": "東京都区部",
    }
    for column in metrics:
        values = pd.to_numeric(wards[column], errors="coerce")
        record[column] = values.sum(min_count=1)
    canonical["13000"] = {
        "municipality_id": "R13000",
        "municipality_code": "13000",
        "prefecture": "東京都",
        "name": "東京都区部",
    }
    return record


def load_main_offices(
    codes: set[str],
    names: dict[str, str],
) -> dict[str, dict[str, Any]]:
    result: dict[str, dict[str, Any]] = {}
    expected: dict[tuple[str, str], list[str]] = {}
    for code, name in names.items():
        if code == "13000":
            continue
        expected.setdefault((code[:2], name + "役所"), []).append(code)
    with zipfile.ZipFile(P05_ARCHIVE) as outer:
        prefixes = sorted({code[:2] for code in codes if code != "13000"})
        for prefix in prefixes:
            inner_name = f"P05-22_{prefix}_GML.zip"
            with zipfile.ZipFile(io.BytesIO(outer.read(inner_name))) as inner:
                payload = json.loads(inner.read(f"P05-22_{prefix}.geojson"))
            for feature in payload["features"]:
                properties = feature["properties"]
                admin_code = str(properties["P05_001"])
                facility_name = str(properties["P05_003"])
                matches = expected.get((admin_code[:2], facility_name), [])
                if str(properties["P05_002"]) != "1" or not matches:
                    continue
                longitude, latitude = feature["geometry"]["coordinates"]
                for code in matches:
                    result[code] = {
                        "facility_admin_code": admin_code,
                        "facility_class_code": str(properties["P05_002"]),
                        "representative_point_name": facility_name,
                        "representative_point_address": str(properties.get("P05_004") or ""),
                        "latitude": float(latitude),
                        "longitude": float(longitude),
                        "coordinate_role": "municipal_main_office",
                        "coordinate_reference_year": 2022,
                        "coordinate_source_url": (
                            "https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-P05-v3_0.html"
                        ),
                    }
    return result


def load_airports() -> pd.DataFrame:
    with zipfile.ZipFile(C28_ARCHIVE) as archive:
        airport_data = json.loads(archive.read("UTF-8/C28-21_Airport.geojson"))
        point_data = json.loads(archive.read("UTF-8/C28-21_AirportReferencePoint.geojson"))
    points = {
        str(feature["properties"]["C28_000"]): feature["geometry"]["coordinates"]
        for feature in point_data["features"]
    }
    rows: list[dict[str, Any]] = []
    for feature in airport_data["features"]:
        properties = feature["properties"]
        regular_code = int(properties["C28_011"])
        if properties["C28_004"] != "供用中" or regular_code not in {1, 2}:
            continue
        reference = str(properties["C28_101"]).lstrip("#")
        longitude, latitude = points[reference]
        rows.append(
            {
                "airport_id": f"AP_{reference}",
                "name": str(properties["C28_005"]),
                "latitude": float(latitude),
                "longitude": float(longitude),
                "airport_type_code": int(properties["C28_003"]),
                "airport_type": str(properties["C28_003"]),
                "scheduled_service": True,
                "scheduled_service_code": regular_code,
                "airport_scale": 1.0,
                "airport_scale_label": "national_phase2_distance_only",
                "source_url": "https://nlftp.mlit.go.jp/ksj/gml/datalist/KsjTmplt-C28-2021.html",
                "c28_reference_id": reference,
            }
        )
    return (
        pd.DataFrame(rows)
        .drop_duplicates("airport_id", keep="first")
        .sort_values("airport_id")
        .reset_index(drop=True)
    )


def build_municipalities(
    data: pd.DataFrame,
    canonical: dict[str, dict[str, Any]],
    metrics: list[str],
    search: pd.DataFrame,
    role_config: dict[str, Any],
    phase1_module: Any,
) -> tuple[pd.DataFrame, dict[str, dict[str, Any]]]:
    tokyo_record = aggregate_tokyo_wards(data, canonical, metrics)
    ward_mask = (
        data["都道府県"].eq("東京都")
        & data["地域コード"].str[1:].str.startswith("131")
        & data["市区町村"].str.endswith("区")
    )
    non_ward = data.loc[~ward_mask].copy()
    rows = non_ward.to_dict(orient="records") + [tokyo_record]
    municipalities = pd.DataFrame(rows)
    municipalities["municipality_id"] = municipalities["地域コード"].str.strip()
    municipalities["municipality_code"] = municipalities["municipality_id"].str[1:]
    municipalities["prefecture"] = municipalities["都道府県"].str.strip()
    municipalities["name"] = municipalities["市区町村"].str.strip()
    numeric_metrics = []
    for column in metrics:
        municipalities[column] = pd.to_numeric(municipalities[column], errors="coerce")
        numeric_metrics.append(column)
    offices = load_main_offices(
        set(municipalities["municipality_code"]),
        dict(zip(municipalities["municipality_code"], municipalities["name"], strict=True)),
    )
    office_rows = pd.DataFrame.from_dict(offices, orient="index")
    office_rows.index.name = "municipality_code"
    municipalities = municipalities.merge(
        office_rows.reset_index(),
        on="municipality_code",
        how="left",
        validate="one_to_one",
    )
    for code, override in role_config.get("representative_point_overrides", {}).items():
        override_mask = municipalities["municipality_code"].eq(str(code))
        for column in (
            "representative_point_name",
            "latitude",
            "longitude",
            "coordinate_role",
            "coordinate_reference_year",
            "coordinate_source_url",
        ):
            if column in override:
                municipalities.loc[override_mask, column] = override[column]
    synthetic = role_config["synthetic_cities"]["東京都区部"]
    synthetic_mask = municipalities["municipality_code"].eq("13000")
    for column in ("latitude", "longitude"):
        municipalities.loc[synthetic_mask, column] = synthetic[column]
    municipalities.loc[synthetic_mask, "representative_point_name"] = "東京都区部（合成代表点）"
    municipalities.loc[synthetic_mask, "coordinate_role"] = "synthetic_reviewed_reference_point"
    municipalities.loc[synthetic_mask, "coordinate_reference_year"] = 2022
    municipalities.loc[synthetic_mask, "coordinate_source_url"] = (
        "https://www.metro.tokyo.lg.jp/tosei/tokyoto/profile/gaiyou.html"
    )
    municipalities = phase1_module.add_similarity_features(municipalities)
    search_lookup = search.set_index("municipality_code").to_dict(orient="index")
    municipalities["search_universe"] = municipalities["municipality_code"].isin(search_lookup)
    municipalities["region"] = municipalities["municipality_code"].map(
        {code: item["region"] for code, item in search_lookup.items()}
    )
    municipalities["prefectural_capital"] = municipalities["municipality_code"].map(
        {code: item["prefectural_capital"] for code, item in search_lookup.items()}
    ).fillna(False)
    municipalities["designated_city"] = municipalities["municipality_code"].map(
        {code: item["designated_city"] for code, item in search_lookup.items()}
    ).fillna(False)
    municipalities["core_city"] = municipalities["municipality_code"].map(
        {code: item["core_city"] for code, item in search_lookup.items()}
    ).fillna(False)
    municipalities["phase1_member"] = municipalities["municipality_code"].map(
        {code: item["phase1_member"] for code, item in search_lookup.items()}
    ).fillna(False)
    required_coordinates = municipalities.loc[
        municipalities["search_universe"], ["latitude", "longitude"]
    ].notna().all(axis=1)
    if not required_coordinates.all():
        missing = municipalities.loc[
            municipalities["search_universe"] & ~required_coordinates, "name"
        ].tolist()
        raise ValueError(f"search cities missing representative points: {missing}")
    municipalities = municipalities.sort_values("municipality_code").reset_index(drop=True)
    canonical_normalized = {
        row["municipality_code"]: {
            "municipality_id": row["municipality_id"],
            "name": row["name"],
            "prefecture": row["prefecture"],
        }
        for row in municipalities[
            ["municipality_code", "municipality_id", "name", "prefecture"]
        ].to_dict(orient="records")
    }
    return municipalities, canonical_normalized


def parse_national_flows(
    canonical: dict[str, dict[str, Any]],
) -> tuple[pd.DataFrame, dict[str, Any]]:
    files = sorted(OD_DIR.glob("table3_*.xlsx"))
    if len(files) != 47:
        raise FileNotFoundError(f"expected 47 national OD files; found {len(files)}")
    rows: list[dict[str, Any]] = []
    scanned_rows = 0
    positive_rows = 0
    id_to_stat = {code: stat_id for code, _pref, stat_id in PREFECTURE_FILES}
    for path in files:
        prefix = path.stem.removeprefix("table3_")
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        for row in sheet.iter_rows(min_row=11, max_col=14, values_only=True):
            scanned_rows += 1
            if row[0] != "0_総数":
                continue
            source_code = _normalize_tokyo_code(_area_code(row[3]))
            target_code = _normalize_tokyo_code(_area_code(row[8]))
            if source_code not in canonical or target_code not in canonical:
                continue
            commuters = _numeric(row[11])
            students = _numeric(row[13])
            raw_total = _numeric(row[9])
            if commuters is None or students is None or raw_total is None:
                continue
            total = commuters + students
            if source_code == target_code or total <= 0:
                continue
            rows.append(
                {
                    "source_code": source_code,
                    "target_code": target_code,
                    "commuters": commuters,
                    "students_15plus": _numeric(row[12]) or 0,
                    "students": students,
                    "total": total,
                    "raw_total": raw_total,
                    "year": 2020,
                    "source_stat_inf_id": id_to_stat[prefix],
                    "source_url": (
                        "https://www.e-stat.go.jp/stat-search/database?"
                        "statdisp_id=0003454524"
                    ),
                }
            )
            positive_rows += 1
        workbook.close()
    raw = pd.DataFrame(rows)
    if raw.empty:
        raise ValueError("national Census OD produced no positive rows")
    numeric = [
        "commuters",
        "students_15plus",
        "students",
        "total",
        "raw_total",
    ]
    grouped = raw.groupby(
        ["source_code", "target_code", "year", "source_stat_inf_id", "source_url"],
        as_index=False,
    )[numeric].sum()
    grouped["source"] = "R" + grouped["source_code"]
    grouped["target"] = "R" + grouped["target_code"]
    grouped["source_name"] = grouped["source_code"].map(
        lambda code: canonical[code]["name"]
    )
    grouped["target_name"] = grouped["target_code"].map(
        lambda code: canonical[code]["name"]
    )
    grouped["source_prefecture"] = grouped["source_code"].map(
        lambda code: canonical[code]["prefecture"]
    )
    grouped["target_prefecture"] = grouped["target_code"].map(
        lambda code: canonical[code]["prefecture"]
    )
    grouped["cross_prefecture"] = (
        grouped["source_prefecture"] != grouped["target_prefecture"]
    )
    columns = [
        "source",
        "target",
        "source_code",
        "target_code",
        "source_name",
        "target_name",
        "source_prefecture",
        "target_prefecture",
        "commuters",
        "students_15plus",
        "students",
        "total",
        "raw_total",
        "year",
        "cross_prefecture",
        "source_stat_inf_id",
        "source_url",
    ]
    return grouped[columns].sort_values(["source", "target"]).reset_index(drop=True), {
        "workbook_count": len(files),
        "scanned_rows": scanned_rows,
        "positive_rows_before_tokyo_aggregation": positive_rows,
        "flow_rows_after_aggregation": len(grouped),
        "cross_prefecture_rows": int(grouped["cross_prefecture"].sum()),
        "self_loops_excluded": True,
        "year": 2020,
    }


def build_airport_access(
    municipalities: pd.DataFrame,
    search: pd.DataFrame,
    airports: pd.DataFrame,
) -> pd.DataFrame:
    search_ids = set(search["municipality_id"])
    rows: list[dict[str, Any]] = []
    airports_list = airports.to_dict(orient="records")
    for city in municipalities.loc[municipalities["search_universe"]].to_dict(orient="records"):
        choices = []
        for airport in airports_list:
            distance = _haversine_km(
                float(city["latitude"]),
                float(city["longitude"]),
                float(airport["latitude"]),
                float(airport["longitude"]),
            )
            choices.append((distance, airport))
        distance, airport = min(
            choices,
            key=lambda item: (item[0], item[1]["airport_id"]),
        )
        rows.append(
            {
                "source": city["municipality_id"],
                "target": airport["airport_id"],
                "source_name": city["name"],
                "target_name": airport["name"],
                "distance_km": distance,
                "distance_method": "Haversine_WGS84_mean_earth_radius",
                "rank": 1,
                "airport_scale": airport["airport_scale"],
                "airport_scale_label": airport["airport_scale_label"],
                "municipality_coordinate_role": city["coordinate_role"],
                "airport_coordinate_role": "official_airport_reference_point",
            }
        )
    if set(row["source"] for row in rows) != search_ids:
        raise ValueError("airport access does not cover exactly the 100-city search universe")
    return pd.DataFrame(rows).sort_values("source").reset_index(drop=True)


def quality_report(
    municipalities: pd.DataFrame,
    search: pd.DataFrame,
    flows: pd.DataFrame,
    airports: pd.DataFrame,
    airport_access: pd.DataFrame,
    graph: nx.MultiDiGraph,
    od_diagnostics: dict[str, Any],
) -> dict[str, Any]:
    from app.graph import graph_facts

    checks = []

    def check(name: str, condition: bool, observed: Any, expected: Any) -> None:
        checks.append(
            {
                "name": name,
                "status": "pass" if condition else "fail",
                "severity": "critical",
                "observed": observed,
                "expected": expected,
            }
        )

    check(
        "analysis_municipality_coverage",
        len(municipalities) >= 1700,
        len(municipalities),
        ">= 1700 including Tokyo ward aggregation",
    )
    check("search_municipality_count", len(search) == 100, len(search), 100)
    check(
        "search_code_unique",
        search["municipality_code"].is_unique,
        int(search["municipality_code"].nunique()),
        100,
    )
    check(
        "flow_endpoint_integrity",
        set(flows["source"]).union(flows["target"]).issubset(
            set(municipalities["municipality_id"])
        ),
        len(set(flows["source"]).union(flows["target"])),
        "all endpoints in national analysis universe",
    )
    check(
        "flow_no_self_loops",
        bool((flows["source"] != flows["target"]).all()),
        bool((flows["source"] != flows["target"]).all()),
        True,
    )
    check(
        "flow_total_reconciliation",
        bool((flows["commuters"] + flows["students"] == flows["total"]).all()),
        int((flows["commuters"] + flows["students"] != flows["total"]).sum()),
        0,
    )
    check(
        "cross_prefecture_flows_retained",
        od_diagnostics["cross_prefecture_rows"] > 0,
        od_diagnostics["cross_prefecture_rows"],
        "> 0",
    )
    check(
        "airport_access_coverage",
        len(airport_access) == 100 and airport_access["source"].is_unique,
        len(airport_access),
        100,
    )
    facts = graph_facts(graph)
    check(
        "graph_is_layered_multidigraph",
        facts["multigraph"] and facts["directed"],
        facts,
        "directed MultiDiGraph",
    )
    check(
        "search_flag_coverage",
        int(municipalities["search_universe"].sum()) == 100,
        int(municipalities["search_universe"].sum()),
        100,
    )
    failures = [item for item in checks if item["status"] == "fail"]
    return {
        "generated_at": utc_now(),
        "overall_status": "pass" if not failures else "fail",
        "analysis_universe": {
            "label": "national_municipalities_with_tokyo_ward_aggregation",
            "count": len(municipalities),
        },
        "search_universe": {
            "label": "approved_national_major_100",
            "count": len(search),
        },
        "checks": checks,
        "failed_check_count": len(failures),
        "od_diagnostics": od_diagnostics,
        "graph_facts": facts,
        "limitations": [
            "Tokyo 23 wards are represented by synthetic municipality R13000.",
            (
                "National full-graph centrality is calculated on the analysis universe; "
                "displayed rankings are projected to 100 search cities."
            ),
            "Performance and memory benchmarks are deferred to a later Phase 2 gate.",
        ],
    }


def _project_rankings(result: dict[str, Any], search_ids: set[str]) -> dict[str, Any]:
    projected = dict(result)
    rankings = {}
    for metric, rows in result.get("rankings", {}).items():
        selected = [
            {**row, "search_rank": rank}
            for rank, row in enumerate(
                [row for row in rows if row.get("municipality_id") in search_ids],
                start=1,
            )
        ]
        rankings[metric] = selected
    projected["rankings"] = rankings
    projected["search_universe_projection"] = "100 approved search cities"
    return projected


def run(
    raw_dir: Path = RAW,
    output_dir: Path = PROCESSED,
) -> dict[str, Any]:
    del raw_dir  # Phase 2 paths are fixed to keep provenance paths stable.
    from app.analysis import (
        city_similarity,
        compute_flow_centrality,
        detect_flow_communities,
        find_similar_cities_with_better_airport,
        rank_airport_access,
        strongest_bidirectional_flow,
    )
    from app.graph import build_mobility_graph, graph_facts

    phase1_module = _load_phase1_module()
    search, search_config = load_city_config()
    role_config = yaml.safe_load(ROLE_CONFIG.read_text(encoding="utf-8"))
    data, canonical_raw, metrics = load_ssdse()
    municipalities, canonical = build_municipalities(
        data,
        canonical_raw,
        metrics,
        search,
        role_config,
        phase1_module,
    )
    flows, od_diagnostics = parse_national_flows(canonical)
    commute = flows.loc[flows["commuters"] > 0].copy()
    study = flows.loc[flows["students"] > 0].copy()
    airports = load_airports()
    airport_access = build_airport_access(municipalities, search, airports)
    graph = build_mobility_graph(
        municipalities,
        flow_edges=flows,
        commute_edges=commute,
        study_edges=study,
        airports=airports,
        airport_access=airport_access,
    )
    quality = quality_report(
        municipalities,
        search,
        flows,
        airports,
        airport_access,
        graph,
        od_diagnostics,
    )
    output_dir.mkdir(parents=True, exist_ok=True)
    write_json(output_dir / "quality_report.json", quality)
    if quality["overall_status"] != "pass":
        failed = [item["name"] for item in quality["checks"] if item["status"] == "fail"]
        raise ValueError(f"Phase 2 data quality checks failed: {', '.join(failed)}")

    tables = {
        "municipalities.parquet": municipalities,
        "commute_edges.parquet": commute,
        "study_edges.parquet": study,
        "flow_edges.parquet": flows,
        "airports.parquet": airports,
        "airport_access.parquet": airport_access,
    }
    for filename, frame in tables.items():
        frame.to_parquet(output_dir / filename, index=False)
        print(f"[write] {output_dir / filename} ({len(frame)} rows)")
    search.to_csv(output_dir / "search_universe.csv", index=False)
    write_json(
        output_dir / "search_universe.json",
        {
            "label": "approved_national_major_100",
            "count": len(search),
            "analysis_universe_label": (
                "national_municipalities_with_tokyo_ward_aggregation"
            ),
            "municipality_ids": search["municipality_id"].tolist(),
        },
    )
    write_json(
        output_dir / "universe_metadata.json",
        {
            "analysis_universe": "national_municipalities_with_tokyo_ward_aggregation",
            "analysis_municipality_count": len(municipalities),
            "search_universe": "approved_national_major_100",
            "search_municipality_count": len(search),
            "phase1_compatibility": "phase1/ remains independent and unchanged",
            "synthetic_entities": {"R13000": "東京都区部"},
        },
    )
    flows.to_csv(output_dir / "flow_edges.csv", index=False)

    specs, population_column = phase1_module.load_feature_specs()
    search_municipalities = municipalities.loc[municipalities["search_universe"]].copy()
    search_ids = set(search_municipalities["municipality_id"])
    centrality = compute_flow_centrality(graph)
    centrality["search_projection"] = _project_rankings(centrality, search_ids)
    metrics_rows = []
    for row in centrality["cities"]:
        metrics_rows.append(
            {
                "municipality_id": row["municipality_id"],
                "name": row["name"],
                "prefecture": row["prefecture"],
                **row["metrics"],
                **{f"rank_{key}": value for key, value in row["ranks"].items()},
                "search_universe": row["municipality_id"] in search_ids,
            }
        )
    pd.DataFrame(metrics_rows).to_parquet(
        output_dir / "national_flow_metrics.parquet",
        index=False,
    )
    hiroshima_id = "R34100"
    overall_similarity = city_similarity(
        search_municipalities,
        hiroshima_id,
        feature_specs=specs,
        population_column=population_column,
        top_n=10,
    )
    medical_specs = [spec for spec in specs if spec.category == "medical"]
    medical_similarity = city_similarity(
        search_municipalities,
        hiroshima_id,
        feature_specs=medical_specs,
        population_column=population_column,
        top_n=10,
    )
    airport_rank = rank_airport_access(
        graph,
        municipality_ids=search_ids,
        max_access_rank=1,
    )
    analyses = {
        "generated_at": utc_now(),
        "analysis_universe": "national_municipalities_with_tokyo_ward_aggregation",
        "search_universe": "approved_national_major_100",
        "strongest_flow_from_hiroshima": _project_rankings(
            strongest_bidirectional_flow(graph, hiroshima_id),
            search_ids,
        ),
        "flow_hubs": centrality["search_projection"],
        "flow_communities": detect_flow_communities(
            graph,
            top_cross_prefecture_links=20,
        ),
        "overall_similarity_to_hiroshima": overall_similarity,
        "medical_similarity_to_hiroshima": medical_similarity,
        "airport_access": airport_rank,
        "similar_with_better_airport": find_similar_cities_with_better_airport(
            search_municipalities,
            graph,
            hiroshima_id,
            feature_specs=specs,
            population_column=population_column,
            max_access_rank=1,
            require_nearer=True,
            top_n=10,
        ),
    }
    write_json(output_dir / "phase2_analysis.json", analyses)

    search_nodes = set(search_ids)
    for city_id in search_ids:
        prefecture = graph.nodes[city_id].get("prefecture")
        if prefecture in graph:
            search_nodes.add(prefecture)
    search_nodes.update(airport_access["target"].tolist())
    search_graph = graph.subgraph(search_nodes).copy()
    write_json(output_dir / "knowledge_graph.json", nx.node_link_data(search_graph))
    write_json(
        output_dir / "analysis_graph_facts.json",
        {
            "full_analysis_graph": graph_facts(graph),
            "search_graph": graph_facts(search_graph),
            "note": (
                "knowledge_graph.json is the 100-city search projection; "
                "Parquet flow tables retain the national analysis universe."
            ),
        },
    )
    report = {
        "generated_at": utc_now(),
        "pipeline": "pipeline/06_prepare_phase2.py",
        "phase": 2,
        "search_city_count": len(search),
        "analysis_municipality_count": len(municipalities),
        "outputs": {
            filename: {"rows": len(frame), "columns": len(frame.columns)}
            for filename, frame in tables.items()
        },
        "national_flow_metrics": {
            "rows": len(metrics_rows),
            "columns": len(metrics_rows[0]) if metrics_rows else 0,
        },
        "quality_report": "data/processed/quality_report.json",
        "analysis_snapshot": "data/processed/phase2_analysis.json",
        "wikipedia_used": False,
        "graphrag_used": False,
        "llm_generated_facts": False,
        "performance_benchmark_status": "deferred",
        "source_selection_report": "docs/city_selection.md",
    }
    write_json(output_dir / "processing_report.json", report)
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--raw-dir", type=Path, default=RAW)
    parser.add_argument("--output-dir", type=Path, default=PROCESSED)
    args = parser.parse_args()
    report = run(args.raw_dir, args.output_dir)
    print(json.dumps(report, ensure_ascii=False, indent=2, default=_json_default))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
