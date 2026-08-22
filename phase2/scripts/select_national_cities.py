"""Build a provisional, reproducible national-major-city selection report.

The output is a review artifact, not the final Phase 2 production city set.
Selection uses mandatory prefectural-capital/designated-city anchors and
transparent regional quotas. National Census OD totals are used as a
selection diagnostic; the full Phase 2 analysis pipeline is implemented only
after the city list is reviewed.
"""

from __future__ import annotations

import argparse
import io
import json
import math
import re
import zipfile
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import pandas as pd
import yaml
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data" / "raw"
PROCESSED = ROOT / "data" / "processed"
ROLE_CONFIG = ROOT / "configs" / "phase2_city_roles.yml"
PHASE1_CONFIG = ROOT / "configs" / "phase1_cities.yml"
FINAL_CITY_CONFIG = ROOT / "configs" / "phase2_cities.yml"
OD_DIR = RAW / "census_od_national"
P05_ARCHIVE = RAW / "geography" / "P05-22_GML.zip"
C28_ARCHIVE = RAW / "airports" / "C28-21_GML.zip"


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _json_default(value: Any) -> Any:
    if value is pd.NA:
        return None
    if hasattr(value, "item"):
        return value.item()
    raise TypeError(f"cannot serialize {type(value).__name__}")


def _write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, default=_json_default) + "\n",
        encoding="utf-8",
    )


def _numeric(value: Any) -> int | None:
    if value is None or str(value).strip() in {"", "-", "—", "–"}:
        return 0
    try:
        return int(float(value))
    except (TypeError, ValueError):
        return None


def _code(value: Any) -> str | None:
    match = re.match(r"^(\d{5})_", str(value or "").strip())
    return match.group(1) if match else None


def _normalize_tokyo_code(code: str | None) -> str | None:
    if code and code.startswith("131"):
        return "13000"
    return code


def _load_config() -> dict[str, Any]:
    return yaml.safe_load(ROLE_CONFIG.read_text(encoding="utf-8"))


def _load_ssdse() -> tuple[pd.DataFrame, dict[str, dict[str, str]]]:
    path = RAW / "SSDSE-A-2026.csv"
    raw: pd.DataFrame | None = None
    for encoding in ("utf-8-sig", "cp932", "euc_jp"):
        try:
            raw = pd.read_csv(path, header=None, dtype=str, encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    if raw is None:
        raise UnicodeError(f"cannot decode {path}")
    data = raw.iloc[3:].reset_index(drop=True)
    data.columns = raw.iloc[2].fillna("").astype(str).tolist()
    records = data[["地域コード", "都道府県", "市区町村", "総人口"]].copy()
    for column in ("地域コード", "都道府県", "市区町村"):
        records[column] = records[column].astype(str).str.strip()
    records["municipality_code"] = records["地域コード"].str.removeprefix("R")
    records["population"] = pd.to_numeric(records["総人口"], errors="coerce").fillna(0).astype(int)
    records = records[records["municipality_code"].str.fullmatch(r"\d{5}")].copy()

    canonical = {
        row["municipality_code"]: {
            "municipality_id": row["地域コード"],
            "prefecture": row["都道府県"],
            "name": row["市区町村"],
        }
        for row in records.to_dict(orient="records")
    }
    return records, canonical


def _add_tokyo_ward_entity(
    records: pd.DataFrame,
    canonical: dict[str, dict[str, str]],
) -> pd.DataFrame:
    wards = records[
        (records["都道府県"] == "東京都")
        & records["municipality_code"].str.startswith("131")
        & records["市区町村"].str.endswith("区")
    ]
    if len(wards) != 23:
        raise ValueError(f"Tokyo ward aggregation expected 23 wards; found {len(wards)}")
    row = pd.DataFrame(
        [
            {
                "地域コード": "R13000",
                "都道府県": "東京都",
                "市区町村": "東京都区部",
                "総人口": int(wards["population"].sum()),
                "municipality_code": "13000",
                "population": int(wards["population"].sum()),
            }
        ]
    )
    canonical["13000"] = {
        "municipality_id": "R13000",
        "prefecture": "東京都",
        "name": "東京都区部",
    }
    return pd.concat([records, row], ignore_index=True)


def _load_phase1_ids() -> set[str]:
    config = yaml.safe_load(PHASE1_CONFIG.read_text(encoding="utf-8"))
    return {str(item["municipality_code"]) for item in config["cities"]}


def _load_od_hubs(
    canonical: dict[str, dict[str, str]],
) -> tuple[pd.DataFrame, dict[str, Any]]:
    files = sorted(OD_DIR.glob("table3_*.xlsx"))
    if len(files) != 47:
        raise FileNotFoundError(
            f"{OD_DIR} must contain 47 national Table 3 workbooks; found {len(files)}. "
            "Run scripts/download_national_selection_od.py first."
        )
    inflow: Counter[str] = Counter()
    outflow: Counter[str] = Counter()
    scanned_rows = 0
    positive_rows = 0
    for path in files:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[workbook.sheetnames[0]]
        for row in sheet.iter_rows(min_row=11, max_col=14, values_only=True):
            scanned_rows += 1
            if row[0] != "0_総数":
                continue
            source = _normalize_tokyo_code(_code(row[3]))
            target = _normalize_tokyo_code(_code(row[8]))
            if source not in canonical or target not in canonical or source == target:
                continue
            commuters = _numeric(row[11])
            students = _numeric(row[13])
            if commuters is None or students is None:
                continue
            total = commuters + students
            if total <= 0:
                continue
            outflow[source] += total
            inflow[target] += total
            positive_rows += 1
        workbook.close()
    rows = [
        {
            "municipality_code": code,
            "od_inflow": int(inflow[code]),
            "od_outflow": int(outflow[code]),
            "od_total_flow": int(inflow[code] + outflow[code]),
        }
        for code in canonical
    ]
    return pd.DataFrame(rows), {
        "workbook_count": len(files),
        "scanned_rows": scanned_rows,
        "positive_rows": positive_rows,
        "metric": "all directed non-self Census commuters plus students",
        "year": 2020,
    }


def _load_airports() -> pd.DataFrame:
    with zipfile.ZipFile(C28_ARCHIVE) as archive:
        airport_data = json.loads(archive.read("UTF-8/C28-21_Airport.geojson"))
        point_data = json.loads(archive.read("UTF-8/C28-21_AirportReferencePoint.geojson"))
    points = {
        str(feature["properties"]["C28_000"]): feature["geometry"]["coordinates"]
        for feature in point_data["features"]
    }
    rows = []
    for feature in airport_data["features"]:
        properties = feature["properties"]
        if properties["C28_004"] != "供用中" or int(properties["C28_011"]) not in {1, 2}:
            continue
        reference = str(properties["C28_101"]).lstrip("#")
        longitude, latitude = points[reference]
        rows.append(
            {
                "airport_id": f"AP_{reference}",
                "name": str(properties["C28_005"]),
                "latitude": float(latitude),
                "longitude": float(longitude),
            }
        )
    return pd.DataFrame(rows)


def _load_main_offices(city_codes: set[str]) -> dict[str, tuple[float, float]]:
    result: dict[str, tuple[float, float]] = {}
    with zipfile.ZipFile(P05_ARCHIVE) as outer:
        for prefix in sorted({code[:2] for code in city_codes if code != "13000"}):
            inner_name = f"P05-22_{prefix}_GML.zip"
            with zipfile.ZipFile(io.BytesIO(outer.read(inner_name))) as inner:
                payload = json.loads(inner.read(f"P05-22_{prefix}.geojson"))
            for feature in payload["features"]:
                properties = feature["properties"]
                if str(properties["P05_002"]) != "1":
                    continue
                admin_code = str(properties["P05_001"])
                if admin_code not in city_codes:
                    continue
                longitude, latitude = feature["geometry"]["coordinates"]
                result[admin_code] = (float(latitude), float(longitude))
    return result


def _haversine_km(
    latitude_a: float,
    longitude_a: float,
    latitude_b: float,
    longitude_b: float,
) -> float:
    radius_km = 6_371.0088
    lat_a, lon_a, lat_b, lon_b = map(
        math.radians, (latitude_a, longitude_a, latitude_b, longitude_b)
    )
    delta_lat = lat_b - lat_a
    delta_lon = lon_b - lon_a
    value = (
        math.sin(delta_lat / 2) ** 2
        + math.cos(lat_a) * math.cos(lat_b) * math.sin(delta_lon / 2) ** 2
    )
    return 2 * radius_km * math.asin(math.sqrt(value))


def _add_airport_access(candidates: pd.DataFrame) -> pd.DataFrame:
    airports = _load_airports()
    offices = _load_main_offices(set(candidates["municipality_code"]))
    airport_rows = airports.to_dict(orient="records")
    distances: list[float | None] = []
    names: list[str | None] = []
    for row in candidates.to_dict(orient="records"):
        point = offices.get(row["municipality_code"])
        if point is None:
            distances.append(None)
            names.append(None)
            continue
        distance, airport = min(
            (
                (
                    _haversine_km(point[0], point[1], item["latitude"], item["longitude"]),
                    item,
                )
                for item in airport_rows
            ),
            key=lambda item: (item[0], item[1]["airport_id"]),
        )
        distances.append(distance)
        names.append(airport["name"])
    result = candidates.copy()
    result["nearest_airport_distance_km"] = distances
    result["nearest_airport"] = names
    result["airport_access_within_50km"] = result["nearest_airport_distance_km"].le(50)
    return result


def _score_and_select(candidates: pd.DataFrame, config: dict[str, Any]) -> pd.DataFrame:
    result = candidates.copy()
    result["population_rank"] = result["population"].rank(method="min", ascending=False).astype(int)
    result["od_hub_rank"] = result["od_total_flow"].rank(method="min", ascending=False).astype(int)
    population_max = max(len(result) - 1, 1)
    result["population_percentile"] = 1 - (result["population_rank"] - 1) / population_max
    result["od_percentile"] = 1 - (result["od_hub_rank"] - 1) / population_max
    result["selection_score"] = (
        0.4 * result["population_percentile"]
        + 0.3 * result["od_percentile"]
        + 0.2 * result["core_city"].astype(int)
        + 0.1 * result["airport_access_within_50km"].astype(int)
    )
    result["provisional_selected"] = False
    selected: set[str] = set(result.loc[result["mandatory_anchor"], "municipality_code"])
    targets = {str(key): int(value) for key, value in config["region_targets"].items()}
    for region, target in targets.items():
        current = result.loc[
            result["municipality_code"].isin(selected) & (result["region"] == region)
        ]
        needed = max(target - len(current), 0)
        pool = result[
            (result["region"] == region) & ~result["municipality_code"].isin(selected)
        ].sort_values(
            ["selection_score", "population", "municipality_code"],
            ascending=[False, False, True],
        )
        selected.update(pool.head(needed)["municipality_code"])
    if len(selected) < 100:
        pool = result[~result["municipality_code"].isin(selected)].sort_values(
            ["selection_score", "population", "municipality_code"],
            ascending=[False, False, True],
        )
        selected.update(pool.head(100 - len(selected))["municipality_code"])
    if len(selected) != 100:
        raise ValueError(f"provisional selection must contain 100 cities; found {len(selected)}")
    result["provisional_selected"] = result["municipality_code"].isin(selected)
    result["provisional_rank"] = (
        result.loc[result["provisional_selected"]]
        .sort_values(
            ["region", "selection_score", "municipality_code"],
            ascending=[True, False, True],
        )
        .groupby("region")
        .cumcount()
        + 1
    )
    return result


def _reason(row: dict[str, Any]) -> str:
    reasons = []
    if row["mandatory_anchor"]:
        reasons.append("県庁所在地/政令指定都市アンカー")
    if row["core_city"]:
        reasons.append("中核市")
    if row["phase1_member"]:
        reasons.append("Phase 1継続候補")
    if row["od_hub_rank"] <= 100:
        reasons.append("全国ODハブ上位100")
    if row["airport_access_within_50km"]:
        reasons.append("定期便空港まで50km以内")
    if not reasons:
        reasons.append("人口規模・地域枠による代表候補")
    return "、".join(reasons)


def write_report(
    candidates: pd.DataFrame,
    od_diagnostics: dict[str, Any],
    config: dict[str, Any],
    path: Path,
    approved: bool,
) -> None:
    selected = candidates[candidates["provisional_selected"]].copy()
    excluded = candidates[~candidates["provisional_selected"]].sort_values(
        ["selection_score", "population"], ascending=[False, False]
    ).head(20)
    selected["reason"] = [_reason(row) for row in selected.to_dict(orient="records")]
    region_summary = (
        selected.groupby("region", sort=False)
        .agg(
            cities=("municipality_code", "count"),
            population=("population", "sum"),
            mandatory=("mandatory_anchor", "sum"),
        )
        .reset_index()
    )
    lines = [
        "# 全国主要100都市候補選定資料（Phase 2 provisional）",
        "",
        f"生成日時: `{utc_now()}`",
        "",
        (
            "**レビュー結果: ユーザー承認済み。`configs/phase2_cities.yml`へ反映済み。**"
            if approved
            else (
                "**重要: この100都市はレビュー前の暫定候補であり、"
                "Phase 2のproduction設定ではありません。**"
            )
        ),
        "",
        "## 選定方針",
        "",
        "- 都道府県庁所在地47都市と政令指定都市20都市の重複を除いたアンカーを必須採用する。",
        (
            "- 東京都の都庁所在地は、SSDSEに単一コードがないため"
            "23区を集約した「東京都区部」を使用する。"
        ),
        (
            "- 残りの枠は地方別枠を守り、人口規模、全国ODハブ性、中核市、"
            "空港アクセスを透明な補助基準にする。"
        ),
        (
            "- 選定スコアは `0.4×人口順位正規化 + 0.3×OD総流入出順位正規化 "
            "+ 0.2×中核市 + 0.1×50km以内空港`。"
        ),
        (
            "- ODハブ性は2020年国勢調査Table 3の全都道府県ODを読み、"
            "非自都市の通勤者＋通学者を全国自治体単位で集計した。"
        ),
        "- 生活圏・PageRank等の全国分析は、都市リストレビュー後の次Stageで実装する。",
        "",
        "## 使用データと状態",
        "",
        (
            f"- SSDSE-A-2026: 全国自治体人口・コード（`{len(candidates) - 1}`"
            "自治体級候補 + 東京都区部集約）"
        ),
        f"- 2020年国勢調査 Table 3: 47都道府県ファイル、{od_diagnostics['scanned_rows']:,}行を走査",
        "- P05-22: 市役所本庁舎代表点",
        "- C28-21: 供用中・定期便対象空港",
        (
            "- 国勢調査ファイルのSHA-256は "
            "`data/provenance/national_selection_od_manifest.json` に記録"
            "（候補選定bootstrap版）"
        ),
        "",
        "## 地方別構成",
        "",
        "| 地方 | 暫定都市数 | 人口合計 | 必須アンカー |",
        "|---|---:|---:|---:|",
    ]
    for row in region_summary.to_dict(orient="records"):
        lines.append(
            f"| {row['region']} | {int(row['cities'])} | {int(row['population']):,} | "
            f"{int(row['mandatory'])} |"
        )
    lines += [
        "",
        "## 暫定採用100都市",
        "",
        (
            "| 地方 | 都道府県 | 都市 | コード | 人口 | 県庁所在地 | 政令市 | "
            "中核市 | OD順位 | 最寄り空港 | 距離km | 採用理由 |"
        ),
        "|---|---|---|---|---:|:---:|:---:|:---:|---:|---|---:|---|",
    ]
    selected_rows = selected.sort_values(
        ["region", "selection_score", "municipality_code"],
        ascending=[True, False, True],
    ).to_dict(orient="records")
    for row in selected_rows:
        distance = (
            "—"
            if pd.isna(row["nearest_airport_distance_km"])
            else f"{row['nearest_airport_distance_km']:.1f}"
        )
        airport_name = (
            "—" if pd.isna(row["nearest_airport"]) else str(row["nearest_airport"])
        )
        lines.append(
            f"| {row['region']} | {row['都道府県']} | {row['name']} | {row['municipality_code']} | "
            f"{int(row['population']):,} | {'○' if row['prefectural_capital'] else ''} | "
            f"{'○' if row['designated_city'] else ''} | {'○' if row['core_city'] else ''} | "
            f"{int(row['od_hub_rank'])} | {airport_name} | {distance} | {row['reason']} |"
        )
    lines += [
        "",
        "## 主要な除外・迷い候補（上位20）",
        "",
        "| 地方 | 都道府県 | 都市 | 人口 | OD順位 | スコア | 除外理由/レビュー事項 |",
        "|---|---|---|---:|---:|---:|---|",
    ]
    for row in excluded.to_dict(orient="records"):
        lines.append(
            f"| {row['region']} | {row['都道府県']} | {row['name']} | {int(row['population']):,} | "
            f"{int(row['od_hub_rank'])} | {row['selection_score']:.4f} | "
            "暫定枠外。地域枠・役割重複・感度分析で再レビュー |"
        )
    lines += [
        "",
        "## 手動レビューが必要な事項",
        "",
        "- 必須アンカー集合と地方別枠の妥当性。",
        "- 東京都区部を合成エンティティとして扱うことの妥当性。",
        "- 中核市リストの基準日（総務省2023-04-01）とPhase 2実装時点の差分。",
        "- 空港50kmを交通拠点性の補助条件に使うことの妥当性。",
        "- ODハブ順位を全国候補選定に使う際の人口規模補正の要否。",
        "- Phase 1の20都市のうち全国100都市から外れる都市の互換性。",
        "",
        "## 承認後の扱い" if approved else "## 次の停止条件",
        "",
        (
            "この資料の100都市をPhase 2のSearch Universeとして固定する。変更する場合は、"
            "選定資料、設定、全国処理済みデータ、Acceptanceを一体で再生成する。"
            if approved
            else (
                "この資料へのレビューが完了するまで、`configs/phase2_cities.yml`は作成せず、"
                "全国本番パイプライン・全国Query・全国Acceptanceを確定扱いにしない。"
            )
        ),
        "",
        "関連する機械可読候補表: `data/processed/national_city_candidates.csv`",
        "",
    ]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(lines), encoding="utf-8")


def run(
    *,
    output_csv: Path = PROCESSED / "national_city_candidates.csv",
    output_json: Path = PROCESSED / "national_city_selection_summary.json",
    output_markdown: Path = ROOT / "docs" / "city_selection.md",
) -> dict[str, Any]:
    config = _load_config()
    records, canonical = _load_ssdse()
    records = _add_tokyo_ward_entity(records, canonical)
    city_records = records[records["市区町村"].str.endswith("市")].copy()
    city_records = pd.concat(
        [city_records, records[records["市区町村"] == "東京都区部"]],
        ignore_index=True,
    ).drop_duplicates("municipality_code")
    city_records["name"] = city_records["市区町村"]
    city_records["region"] = city_records["都道府県"].map(config["regions"])
    if city_records["region"].isna().any():
        raise ValueError("some candidate cities do not have a region mapping")
    capital_map = config["prefectural_capitals"]
    designated = set(config["designated_cities"])
    core = set(config["core_cities"])
    phase1 = _load_phase1_ids()
    city_records["prefectural_capital"] = [
        name == capital_map.get(pref)
        for pref, name in zip(city_records["都道府県"], city_records["name"], strict=True)
    ]
    city_records["designated_city"] = city_records["name"].isin(designated)
    city_records["core_city"] = city_records["name"].isin(core)
    city_records["mandatory_anchor"] = (
        city_records["prefectural_capital"] | city_records["designated_city"]
    )
    city_records["phase1_member"] = city_records["municipality_code"].isin(phase1)
    city_records["city_role"] = city_records.apply(
        lambda row: " / ".join(
            role
            for role, enabled in (
                ("prefectural_capital", row["prefectural_capital"]),
                ("designated_city", row["designated_city"]),
                ("core_city", row["core_city"]),
            )
            if enabled
        )
        or "regional_candidate",
        axis=1,
    )
    od, od_diagnostics = _load_od_hubs(canonical)
    city_records = city_records.merge(od, on="municipality_code", how="left", validate="one_to_one")
    for column in ("od_inflow", "od_outflow", "od_total_flow"):
        city_records[column] = city_records[column].fillna(0).astype(int)
    city_records = _add_airport_access(city_records)
    city_records = _score_and_select(city_records, config)
    city_records["reason"] = [_reason(row) for row in city_records.to_dict(orient="records")]
    approved = False
    if FINAL_CITY_CONFIG.exists():
        final_config = yaml.safe_load(FINAL_CITY_CONFIG.read_text(encoding="utf-8")) or {}
        approved = final_config.get("status") == "approved_by_user"
    city_records.to_csv(output_csv, index=False, encoding="utf-8")
    summary = {
        "generated_at": utc_now(),
        "status": "approved_by_user" if approved else "provisional_review_required",
        "candidate_count": len(city_records),
        "provisional_selected_count": int(city_records["provisional_selected"].sum()),
        "mandatory_anchor_count": int(city_records["mandatory_anchor"].sum()),
        "tokyo_ward_aggregation": "R13000 / 23 wards",
        "od_diagnostics": od_diagnostics,
        "selection_formula": (
            "0.4*population_percentile + 0.3*od_percentile + "
            "0.2*core_city + 0.1*airport_access_within_50km"
        ),
        "output_csv": str(output_csv.relative_to(ROOT)),
        "output_markdown": str(output_markdown.relative_to(ROOT)),
        "final_city_set_confirmed": approved,
    }
    _write_json(output_json, summary)
    write_report(city_records, od_diagnostics, config, output_markdown, approved)
    return summary


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output-csv",
        type=Path,
        default=PROCESSED / "national_city_candidates.csv",
    )
    parser.add_argument(
        "--output-json",
        type=Path,
        default=PROCESSED / "national_city_selection_summary.json",
    )
    parser.add_argument(
        "--output-markdown",
        type=Path,
        default=ROOT / "docs" / "city_selection.md",
    )
    args = parser.parse_args()
    summary = run(
        output_csv=args.output_csv,
        output_json=args.output_json,
        output_markdown=args.output_markdown,
    )
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
