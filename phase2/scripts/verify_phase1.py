#!/usr/bin/env python3
"""Verify the Phase 1 acceptance gate and write reproducible evidence reports."""

from __future__ import annotations

import argparse
import json
import math
import os
import sys
from collections.abc import Callable
from datetime import datetime, timezone
from html import escape
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from app.executor import QueryExecutor, execution_is_finite  # noqa: E402
from app.llm import (  # noqa: E402
    AzureOpenAIChatClient,
    AzureOpenAIPlanner,
    AzureOpenAIResponse,
)
from app.planner import (  # noqa: E402
    INTENT_AIRPORT_BETTER,
    INTENT_BIDIRECTIONAL_FLOW,
    INTENT_CATEGORY_SIMILARITY,
    INTENT_CITY_SIMILARITY,
    INTENT_CROSS_PREFECTURE_COMMUNITIES,
    INTENT_FLOW_HUB,
    INTENT_SIMILAR_AIRPORT,
)
from app.response import attach_response_supplement  # noqa: E402
from scripts.download_data import DOWNLOADS, EXPECTED_SHA256, sha256  # noqa: E402

DEFAULT_PROCESSED = ROOT / "data" / "processed"
DEFAULT_JSON_REPORT = DEFAULT_PROCESSED / "acceptance_results.json"
DEFAULT_HTML_REPORT = ROOT / "docs" / "phase1_acceptance.html"
DEFAULT_SIMILARITY_CONFIG = ROOT / "configs" / "similarity_features.yml"


class AcceptanceFailure(AssertionError):
    """A human-readable failed acceptance assertion."""


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AcceptanceFailure(message)


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def _common(result: dict[str, Any], intent: str) -> dict[str, Any]:
    require(result.get("status") == "ok", f"status={result.get('status')!r}, expected 'ok'")
    require(result.get("plan", {}).get("intent") == intent, "unexpected query intent")
    facts = result.get("hard_facts")
    require(isinstance(facts, dict), "hard_facts is missing")
    metadata = facts.get("evidence_metadata")
    require(isinstance(metadata, dict), "evidence_metadata is missing")
    require(bool(metadata.get("datasets")), "dataset editions are missing")
    require(bool(metadata.get("metrics")), "metric definitions/units are missing")
    require(bool(metadata.get("limitations")), "limitations are missing")
    universe = metadata.get("universe", {})
    require(universe.get("actual_municipality_count") == 20, "query universe is not 20 cities")
    boundary = result.get("provenance_boundary", {})
    require(boundary.get("llm_used_for_numbers") is False, "LLM was used for numbers")
    require(
        boundary.get("wikipedia_or_graphrag_used") is False,
        "Wikipedia or GraphRAG entered the hard-fact path",
    )
    require(
        boundary.get("feature_config") == "configs/similarity_features.yml",
        "feature-config provenance is missing or not project-relative",
    )
    require(execution_is_finite(result), "result contains a non-finite number")
    json.dumps(result, ensure_ascii=False, allow_nan=False)
    return facts


def _flow(result: dict[str, Any]) -> dict[str, Any]:
    facts = _common(result, INTENT_BIDIRECTIONAL_FLOW)
    strongest = facts.get("strongest") or {}
    require(strongest.get("name") == "廿日市市", "strongest counterpart is not 廿日市市")
    require(strongest.get("forward_flow") == 13_331, "forward flow does not match raw Census")
    require(strongest.get("reverse_flow") == 20_804, "reverse flow does not match raw Census")
    require(strongest.get("bidirectional_flow") == 34_135, "bidirectional sum is not 34,135")
    return {
        "top_city": strongest["name"],
        "hiroshima_to_city_people": strongest["forward_flow"],
        "city_to_hiroshima_people": strongest["reverse_flow"],
        "bidirectional_people": strongest["bidirectional_flow"],
    }


def _hub(result: dict[str, Any]) -> dict[str, Any]:
    facts = _common(result, INTENT_FLOW_HUB)
    rankings = facts.get("rankings", {})
    required = {
        "weighted_in_degree",
        "weighted_out_degree",
        "weighted_degree",
        "weighted_pagerank",
    }
    require(required.issubset(rankings), "one or more required hub metrics are missing")
    leaders = {metric: rows[0] for metric, rows in rankings.items() if metric in required}
    require(
        all(item.get("name") == "広島市" for item in leaders.values()),
        "Hiroshima is not first on every required hub metric",
    )
    return {
        metric: {"city": item["name"], "value": item["value"]}
        for metric, item in sorted(leaders.items())
    }


def _communities(result: dict[str, Any]) -> dict[str, Any]:
    facts = _common(result, INTENT_CROSS_PREFECTURE_COMMUNITIES)
    communities = facts.get("cross_prefecture_communities") or []
    require(bool(communities), "no cross-prefecture mobility community was detected")
    for community in communities:
        require(len(community.get("prefectures", [])) >= 2, "community is not cross-prefecture")
    return {
        "cross_prefecture_community_count": len(communities),
        "communities": [
            {
                "community_id": item.get("community_id"),
                "prefectures": item.get("prefectures"),
                "municipalities": [city.get("name") for city in item.get("municipalities", [])],
                "internal_symmetric_weight_people": item.get("internal_symmetric_weight"),
            }
            for item in communities
        ],
    }


def _similarity(
    result: dict[str, Any],
    *,
    intent: str,
    category: str,
) -> dict[str, Any]:
    facts = _common(result, intent)
    require(facts.get("requested_category") == category, "wrong similarity category")
    ranking = facts.get("ranking") or []
    require(bool(ranking), "similarity ranking is empty")
    require(ranking[0].get("name") == "岡山市", "top similarity candidate is not 岡山市")
    return {
        "category": category,
        "top_city": ranking[0]["name"],
        "cosine_similarity": ranking[0]["similarity"],
        "feature_count": facts.get("feature_count"),
    }


def _overall_similarity(result: dict[str, Any]) -> dict[str, Any]:
    return _similarity(result, intent=INTENT_CITY_SIMILARITY, category="overall")


def _medical_similarity(result: dict[str, Any]) -> dict[str, Any]:
    return _similarity(result, intent=INTENT_CATEGORY_SIMILARITY, category="medical")


def _airport_better(result: dict[str, Any]) -> dict[str, Any]:
    facts = _common(result, INTENT_AIRPORT_BETTER)
    require(facts.get("requested_metric") == "nearest_distance", "generic access used a score")
    reference = facts.get("reference") or {}
    reference_distance = reference.get("nearest_airport", {}).get("distance_km")
    ranking = facts.get("ranking") or []
    require(bool(ranking), "no city is closer to an eligible airport than Hiroshima")
    require(ranking[0].get("name") == "岩国市", "nearest qualifying city is not 岩国市")
    distances = [item.get("nearest_airport", {}).get("distance_km") for item in ranking]
    require(
        all(distance is not None and distance < reference_distance for distance in distances),
        "a candidate is not closer than the reference city",
    )
    return {
        "metric": "nearest scheduled-airport straight-line distance",
        "reference_city": reference.get("name"),
        "reference_airport": reference.get("nearest_airport", {}).get("name"),
        "reference_distance_km": reference_distance,
        "top_city": ranking[0]["name"],
        "top_airport": ranking[0]["nearest_airport"]["name"],
        "top_distance_km": ranking[0]["nearest_airport"]["distance_km"],
    }


def _similar_airport(result: dict[str, Any]) -> dict[str, Any]:
    facts = _common(result, INTENT_SIMILAR_AIRPORT)
    require(
        facts.get("requested_airport_metric") == "nearest_distance",
        "generic combined query used an airport score",
    )
    ranking = facts.get("ranking") or []
    require(bool(ranking), "combined similarity/airport ranking is empty")
    require(ranking[0].get("name") == "岡山市", "top combined candidate is not 岡山市")
    require(
        all(item.get("nearest_distance_improvement_km", 0) > 0 for item in ranking),
        "combined ranking contains a city that is not nearer",
    )
    top = ranking[0]
    airport = top.get("candidate_airport_access", {}).get("nearest_airport", {})
    return {
        "top_city": top["name"],
        "cosine_similarity": top["similarity"],
        "nearest_airport": airport.get("name"),
        "distance_km": airport.get("distance_km"),
        "nearer_than_hiroshima_km": top["nearest_distance_improvement_km"],
        "ranking_rule": "filter by shorter airport distance, then rank by SSDSE similarity",
    }


QUERY_CASES: tuple[tuple[str, str, Callable[[dict[str, Any]], dict[str, Any]]], ...] = (
    ("AC-1", "広島市と人の行き来が最も多い都市は？", _flow),
    ("AC-2", "広島周辺で人流のハブになっている都市は？", _hub),
    ("AC-3", "県境をまたぐ実質的な生活圏はある？", _communities),
    ("AC-4", "広島市と都市特性が似ている都市は？", _overall_similarity),
    ("AC-5", "広島市より空港アクセスが良い都市は？", _airport_better),
    ("AC-6", "広島市と似ていて空港アクセスも良い都市は？", _similar_airport),
    ("AC-7", "医療面で広島市と似ている都市は？", _medical_similarity),
)


def verify_source_files() -> dict[str, Any]:
    verified = []
    for key, (_url, path) in DOWNLOADS.items():
        require(path.is_file(), f"missing source file: {path.relative_to(ROOT)}")
        observed = sha256(path)
        require(observed == EXPECTED_SHA256[key], f"checksum mismatch: {path.relative_to(ROOT)}")
        verified.append(
            {
                "key": key,
                "path": str(path.relative_to(ROOT)),
                "sha256": observed,
            }
        )
    return {"verified_file_count": len(verified), "files": verified}


def verify_raw_census_spot_check() -> dict[str, Any]:
    path = ROOT / "data" / "raw" / "census_od" / "table3_hiroshima.xlsx"
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    wanted = {("34100", "34213"), ("34213", "34100")}
    observed: dict[tuple[str, str], dict[str, int]] = {}
    for row in sheet.iter_rows(min_row=11, max_col=14, values_only=True):
        source = str(row[3] or "").split("_", 1)[0]
        target = str(row[8] or "").split("_", 1)[0]
        if (source, target) not in wanted:
            continue
        observed[(source, target)] = {
            "raw_total": int(row[9]),
            "commuters": int(row[11]),
            "students": int(row[13]),
        }
    workbook.close()
    require(set(observed) == wanted, "raw Census spot-check pairs are missing")
    require(observed[("34100", "34213")]["raw_total"] == 13_331, "raw forward total drifted")
    require(observed[("34213", "34100")]["raw_total"] == 20_804, "raw reverse total drifted")
    for values in observed.values():
        require(
            values["commuters"] + values["students"] == values["raw_total"],
            "raw Census total does not reconcile",
        )
    return {
        "source": "2020 Census Table 3 Hiroshima resident-prefecture workbook",
        "hiroshima_to_hatsukaichi": observed[("34100", "34213")],
        "hatsukaichi_to_hiroshima": observed[("34213", "34100")],
        "bidirectional_total": sum(values["raw_total"] for values in observed.values()),
    }


def verify_airport_distance_spot_check(processed_dir: Path) -> dict[str, Any]:
    municipalities = pd.read_parquet(processed_dir / "municipalities.parquet")
    airports = pd.read_parquet(processed_dir / "airports.parquet")
    access = pd.read_parquet(processed_dir / "airport_access.parquet")
    city = municipalities.loc[municipalities["name"] == "広島市"].iloc[0]

    def haversine(airport: pd.Series) -> float:
        radius = 6_371.0088
        lat_a, lon_a, lat_b, lon_b = map(
            math.radians,
            (city["latitude"], city["longitude"], airport["latitude"], airport["longitude"]),
        )
        delta_lat = lat_b - lat_a
        delta_lon = lon_b - lon_a
        term = (
            math.sin(delta_lat / 2) ** 2
            + math.cos(lat_a) * math.cos(lat_b) * math.sin(delta_lon / 2) ** 2
        )
        return 2 * radius * math.asin(math.sqrt(term))

    candidates = [(haversine(row), row) for _, row in airports.iterrows()]
    distance, airport = min(candidates, key=lambda item: (item[0], item[1]["airport_id"]))
    stored = access.loc[access["source"] == city["municipality_id"]].iloc[0]
    require(airport["airport_id"] == stored["target"], "stored nearest airport is not the minimum")
    require(math.isclose(distance, stored["distance_km"], abs_tol=1e-9), "distance mismatch")
    return {
        "municipality_point": city["representative_point_name"],
        "nearest_airport": airport["name"],
        "distance_km": distance,
        "method": "independent Haversine recomputation",
    }


def verify_missing_condition_boundary(
    executor: QueryExecutor,
    *,
    planner_provider: Any = None,
    response_provider: Any = None,
) -> dict[str, Any]:
    result = executor.execute(
        "自然が豊かで温泉もあって空港が近い都市は？",
        planner_provider=planner_provider,
    )
    if response_provider is not None and result.get("hard_facts") is not None:
        result = attach_response_supplement(result, response_provider)
    require(result.get("status") == "partial", "unsupported-condition query is not partial")
    require(result.get("hard_facts", {}).get("analysis") == "airport_access", "airport part failed")
    missing = {item.get("condition") for item in result.get("missing_structured_conditions", [])}
    require(missing == {"nature", "hot_springs"}, "missing conditions were not separated")
    return {
        "supported_part": "nearest scheduled-airport distance",
        "missing_conditions": sorted(missing),
        "llm_used_for_numbers": False,
    }


def _run_check(name: str, operation: Callable[[], Any]) -> dict[str, Any]:
    try:
        evidence = operation()
    except Exception as exc:
        return {"name": name, "status": "fail", "error": f"{type(exc).__name__}: {exc}"}
    return {"name": name, "status": "pass", "evidence": evidence}


def _format_number(value: Any, digits: int = 3) -> str:
    if value is None:
        return "—"
    number = float(value)
    if number.is_integer():
        return f"{int(number):,}"
    return f"{number:,.{digits}f}"


def _html_text(value: Any) -> str:
    return escape(str(value), quote=True)


def _relative_report_link(report_path: Path, target: Path) -> str:
    return Path(os.path.relpath(target, start=report_path.parent)).as_posix()


def write_html(report: dict[str, Any], path: Path) -> None:
    query_by_id = {item["id"]: item for item in report["queries"]}
    flow = query_by_id.get("AC-1", {}).get("evidence", {})
    hub = query_by_id.get("AC-2", {}).get("evidence", {})
    communities = query_by_id.get("AC-3", {}).get("evidence", {})
    overall = query_by_id.get("AC-4", {}).get("evidence", {})
    airport = query_by_id.get("AC-5", {}).get("evidence", {})
    combined = query_by_id.get("AC-6", {}).get("evidence", {})
    medical = query_by_id.get("AC-7", {}).get("evidence", {})
    status = report["overall_status"].upper()
    status_class = "pass" if report["overall_status"] == "pass" else "fail"
    llm_mode = (
        "Azure OpenAIのPlanner/Responseを使用"
        if report.get("llm_mode")
        else "決定的なオフラインPlanner/Responseを使用"
    )
    query_rows = [
        (
            "AC-1",
            flow.get("top_city", "—"),
            f"{_format_number(flow.get('bidirectional_people'))}人 "
            f"({_format_number(flow.get('hiroshima_to_city_people'))} + "
            f"{_format_number(flow.get('city_to_hiroshima_people'))})",
        ),
        (
            "AC-2",
            "重み付きPageRank 1位: "
            + str(hub.get("weighted_pagerank", {}).get("city", "—")),
            _format_number(hub.get("weighted_pagerank", {}).get("value"), 6),
        ),
        (
            "AC-3",
            "県境をまたぐコミュニティ",
            f"{_format_number(communities.get('cross_prefecture_community_count'))}件",
        ),
        (
            "AC-4",
            overall.get("top_city", "—"),
            f"cosine類似度 {_format_number(overall.get('cosine_similarity'), 4)}",
        ),
        (
            "AC-5",
            f"{airport.get('top_city', '—')} → {airport.get('top_airport', '—')}",
            f"{_format_number(airport.get('top_distance_km'), 1)} km; "
            f"広島市 {_format_number(airport.get('reference_distance_km'), 1)} km",
        ),
        (
            "AC-6",
            combined.get("top_city", "—"),
            f"cosine類似度 {_format_number(combined.get('cosine_similarity'), 4)}、"
            f"空港 {_format_number(combined.get('distance_km'), 1)} km",
        ),
        (
            "AC-7",
            medical.get("top_city", "—"),
            f"医療cosine類似度 {_format_number(medical.get('cosine_similarity'), 4)}",
        ),
    ]
    query_html = []
    for case_id, evidence, quantitative in query_rows:
        item = query_by_id[case_id]
        item_status = item["status"]
        query_html.append(
            "<tr>"
            f"<th scope=\"row\">{_html_text(case_id)}</th>"
            f"<td>{_html_text(item['query'])}</td>"
            f"<td><span class=\"status {_html_text(item_status)}\">"
            f"{_html_text(item_status.upper())}</span></td>"
            f"<td>{_html_text(evidence)}: {_html_text(quantitative)}</td>"
            "</tr>"
        )

    check_rows = []
    for item in report["checks"]:
        item_status = item["status"]
        detail = item.get("error", "")
        check_rows.append(
            "<tr>"
            f"<td>{_html_text(item['name'])}</td>"
            f"<td><span class=\"status {_html_text(item_status)}\">"
            f"{_html_text(item_status.upper())}</span></td>"
            f"<td>{_html_text(detail)}</td>"
            "</tr>"
        )

    seed_rows = []
    for item in report.get("seed_review", [])[:5]:
        bidirectional = f"{int(item['bidirectional_flow_with_seed']):,}人"
        strongest = f"{int(item['strongest_seed_bidirectional_flow']):,}人"
        seed_rows.append(
            "<tr>"
            f"<td>{_html_text(item['rank'])}</td>"
            f"<td>{_html_text(item['candidate_prefecture'] + item['candidate_name'])}</td>"
            f"<td>{_html_text(bidirectional)}</td>"
            f"<td>{_html_text(item['strongest_seed_name'])} "
            f"({_html_text(strongest)})</td>"
            "</tr>"
        )

    json_href = _relative_report_link(path, ROOT / "data" / "processed" / "acceptance_results.json")
    provenance_href = _relative_report_link(path, ROOT / "data" / "provenance")
    html_document = f"""<!doctype html>
<html lang="ja">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>Phase 1受入検証レポート</title>
  <style>
    :root {{
      color-scheme: light;
      font-family: -apple-system, BlinkMacSystemFont, "Hiragino Sans", sans-serif;
    }}
    body {{ background: #f4f6f8; color: #1f2933; margin: 0; line-height: 1.6; }}
    main {{ max-width: 1180px; margin: 0 auto; padding: 2rem 1rem 4rem; }}
    header, section {{
      background: #fff;
      border: 1px solid #d9e2ec;
      border-radius: 10px;
      margin: 1rem 0;
      padding: 1.25rem;
    }}
    h1, h2 {{ color: #102a43; margin-top: 0; }}
    .summary {{ align-items: center; display: flex; gap: 1rem; flex-wrap: wrap; }}
    .overall {{ border-radius: 999px; font-weight: 700; padding: .25rem .8rem; }}
    .overall.pass, .status.pass {{ background: #d9f2e6; color: #0b6b3a; }}
    .overall.fail, .status.fail {{ background: #ffe3e3; color: #a61b1b; }}
    .status {{
      border-radius: 999px;
      display: inline-block;
      font-size: .8rem;
      font-weight: 700;
      padding: .1rem .5rem;
    }}
    table {{ border-collapse: collapse; margin: .75rem 0; width: 100%; }}
    th, td {{
      border-bottom: 1px solid #d9e2ec;
      padding: .6rem;
      text-align: left;
      vertical-align: top;
    }}
    th {{ background: #f0f4f8; color: #243b53; }}
    code {{ background: #f0f4f8; border-radius: 4px; padding: .1rem .3rem; }}
    a {{ color: #0b69a3; }}
    ul {{ padding-left: 1.4rem; }}
    .muted {{ color: #52606d; }}
  </style>
</head>
<body>
<main>
  <header>
    <div class="summary">
      <h1>Phase 1受入検証レポート</h1>
      <span class="overall {_html_text(status_class)}">{_html_text(status)}</span>
    </div>
    <p class="muted">生成日時: {_html_text(report['generated_at'])} / scripts/verify_phase1.py</p>
    <p>{_html_text(llm_mode)}。LLMはランキング・数値計算には使用していません。</p>
  </header>

  <section>
    <h2>受入条件</h2>
    <table>
      <thead><tr><th>ID</th><th>自然言語質問</th><th>状態</th><th>定量的根拠</th></tr></thead>
      <tbody>{''.join(query_html)}</tbody>
    </table>
    <p class="muted">空港アクセスは、定期便のある空港へのHaversine直線距離です。
      旅客規模は明示的に要求した場合だけ使用します。</p>
  </section>

  <section>
    <h2>独立チェック</h2>
    <table>
      <thead><tr><th>チェック</th><th>状態</th><th>詳細</th></tr></thead>
      <tbody>{''.join(check_rows)}</tbody>
    </table>
  </section>

  <section>
    <h2>Seedレビュー（提案用の根拠）</h2>
    <p>Phase 1の20都市seedは変更していません。以下は将来レビュー用の候補です。</p>
    <table>
      <thead><tr><th>順位</th><th>候補</th><th>seed全体との双方向人流</th><th>最強のseed接続</th></tr></thead>
      <tbody>{''.join(seed_rows)}</tbody>
    </table>
  </section>

  <section>
    <h2>範囲と制約</h2>
    <ul>
      <li>人流ランキング、PageRank、communityは固定20都市の誘導グラフ内の結果であり、全国ランキングではありません。</li>
      <li>国勢調査ODの自都市ループは都市間edgeから除外しています。</li>
      <li>空港距離は市役所代表点からC28参照点までの直線距離であり、道路・鉄道・ターミナル所要時間ではありません。</li>
      <li>C28空港情報は2021年、SSDSE各項目は異なる参照年のスナップショットです。</li>
      <li>自然、温泉、家賃、気候、観光、災害リスクはPhase 1の構造化事実には含まれません。</li>
    </ul>
  </section>

  <section>
    <h2>証跡</h2>
    <ul>
      <li><a href="{_html_text(json_href)}">機械可読な受入結果（JSON）</a></li>
      <li><a href="{_html_text(provenance_href)}">
        データ出典・ライセンス・変換・チェックサム</a></li>
    </ul>
  </section>
</main>
</body>
</html>
"""
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(html_document, encoding="utf-8")


def run(
    processed_dir: Path = DEFAULT_PROCESSED,
    json_report: Path = DEFAULT_JSON_REPORT,
    html_report: Path = DEFAULT_HTML_REPORT,
    similarity_config: Path = DEFAULT_SIMILARITY_CONFIG,
    use_llm: bool = False,
) -> dict[str, Any]:
    checks = [
        _run_check("10 pinned source SHA-256 digests", verify_source_files),
        _run_check(
            "processed data quality report",
            lambda: _verify_quality_report(processed_dir),
        ),
        _run_check("raw Census OD spot check", verify_raw_census_spot_check),
        _run_check(
            "independent Hiroshima airport-distance recomputation",
            lambda: verify_airport_distance_spot_check(processed_dir),
        ),
    ]
    executor = QueryExecutor.from_processed_dir(
        processed_dir,
        similarity_config=similarity_config,
    )
    planner_provider = None
    response_provider = None
    if use_llm:
        client = AzureOpenAIChatClient.from_environment()
        planner_provider = AzureOpenAIPlanner(client)
        response_provider = AzureOpenAIResponse(client)

    queries = []
    raw_results: dict[str, Any] = {}
    for case_id, query, validator in QUERY_CASES:
        result = executor.execute(query, planner_provider=planner_provider)
        if response_provider is not None and result.get("hard_facts") is not None:
            result = attach_response_supplement(result, response_provider)
        raw_results[case_id] = result
        checked = _run_check(case_id, lambda result=result, validator=validator: validator(result))
        queries.append({"id": case_id, "query": query, **checked})
    checks.append(
        _run_check(
            "unsupported-condition evidence boundary",
            lambda: verify_missing_condition_boundary(
                executor,
                planner_provider=planner_provider,
                response_provider=response_provider,
            ),
        )
    )

    seed_path = processed_dir / "seed_membership_review.csv"
    seed_review = pd.read_csv(seed_path).head(10).to_dict(orient="records")
    failed = [item for item in [*checks, *queries] if item["status"] != "pass"]
    report = {
        "generated_at": utc_now(),
        "overall_status": "pass" if not failed else "fail",
        "phase": 1,
        "scope": "fixed 20-city induced universe",
        "checks": checks,
        "queries": queries,
        "query_results": raw_results,
        "seed_review": seed_review,
        "failed_count": len(failed),
        "pdf_used": False,
        "wikipedia_used": False,
        "graphrag_used": False,
        "llm_used_for_numbers": False,
        "llm_mode": use_llm,
    }
    json_report.parent.mkdir(parents=True, exist_ok=True)
    json_report.write_text(
        json.dumps(report, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )
    write_html(report, html_report)
    return report


def _verify_quality_report(processed_dir: Path) -> dict[str, Any]:
    path = processed_dir / "quality_report.json"
    quality = json.loads(path.read_text(encoding="utf-8"))
    require(quality.get("overall_status") == "pass", "quality_report is not pass")
    require(quality.get("failed_check_count") == 0, "quality_report has failed checks")
    require(len(quality.get("checks", [])) >= 15, "quality_report is incomplete")
    return {
        "overall_status": quality["overall_status"],
        "passed_check_count": len(quality["checks"]),
        "graph_facts": quality.get("graph_facts"),
    }


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--processed-data", type=Path, default=DEFAULT_PROCESSED)
    parser.add_argument("--json-report", type=Path, default=DEFAULT_JSON_REPORT)
    parser.add_argument("--html-report", type=Path, default=DEFAULT_HTML_REPORT)
    parser.add_argument("--similarity-config", type=Path, default=DEFAULT_SIMILARITY_CONFIG)
    parser.add_argument(
        "--llm",
        action="store_true",
        help="Azure OpenAIのPlanner/Response経路も含めて受入検証する。",
    )
    args = parser.parse_args(argv)
    report = run(
        args.processed_data,
        args.json_report,
        args.html_report,
        args.similarity_config,
        args.llm,
    )
    print(
        json.dumps(
            {
                "overall_status": report["overall_status"],
                "passed_queries": sum(item["status"] == "pass" for item in report["queries"]),
                "failed_count": report["failed_count"],
                "llm_mode": report["llm_mode"],
                "json_report": str(args.json_report),
                "html_report": str(args.html_report),
            },
            ensure_ascii=False,
            indent=2,
        )
    )
    return 0 if report["overall_status"] == "pass" else 1


if __name__ == "__main__":
    raise SystemExit(main())
